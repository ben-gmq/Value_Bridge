"""Slice 1e-a: the FUNCTION_REQUIREMENT import target (spec docs/slice1e_imports_spec.md §7.0, §7.2, round-2 N1/N3/N5/N6;
acceptance §10 items 1, 4–7, 10–12). Grade 2: a wrong FR row reaches a client deliverable but is an audited,
correctable update; added link ids are kept for a manual reversal. Every check runs on a real export, edited the way a
consultant would edit it.

Deliberately uncovered (spec §13): files not made by VB beyond the file-level checks; two batches over the same rows at
once (row lock and row_version, proven in 4a); a retired FR type in use (no route retires a company code in a test
project); a hidden FR type (every default type is pinned and cannot be hidden) — a hidden status proves the same
picker rule (D4, R5)."""
import io

import openpyxl
from sqlalchemy import select

from models import BrFunctionRequirement, FunctionRequirement
from tests.test_baseline import freeze, scope
from tests.test_br_fr import lnk, rows as br_rows, set_components, unlnk
from tests.test_bulk_import import batch, codes_of, commit
from tests.test_code_lists import code_id
from tests.test_function_requirement_api import fr, patch
from tests.test_scope_api import API, ed  # noqa: F401 — ed is a fixture
from tests.test_solution_api import link, sol

URL = API + "/projects/{p}/bulk/function-requirements/{what}"
TAG = "sheet:FUNCTION_REQUIREMENT"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export(client, ed) -> bytes:
    r = client.get(URL.format(p=ed["p"], what="export"), headers=ed["h"])
    assert r.status_code == 200, r.text
    return r.content


def edit(body: bytes, fn) -> bytes:
    """`fn(ws, col, rows, add)`: rows by FR number; `add(**cells)` appends a new row."""
    wb = openpyxl.load_workbook(io.BytesIO(body))
    ws = next(s for s in wb.worksheets if s.cell(1, 1).value == TAG)
    col = {ws.cell(1, i).value: i for i in range(2, ws.max_column + 1)}
    rows = {ws.cell(r, col["fr_number"]).value: r for r in range(3, ws.max_row + 1)}

    def add(**cells):
        at = max(ws.max_row, 2) + 1
        for k, v in cells.items():
            ws.cell(at, col[k]).value = v
    fn(ws, col, rows, add)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def put(ws, row, column, value):
    ws.cell(row, column).value = value


def staged(client, ed, body, expect=201):
    r = client.post(URL.format(p=ed["p"], what="validate"), content=body,
                    headers={**ed["h"], "Content-Type": XLSX, "X-VB-File-Name": "frs.xlsx"})
    assert r.status_code == expect, r.text
    return r.json()


def by_row(client, ed, pv) -> list[dict]:
    return batch(client, ed, pv)["rows"]


def by_key(client, ed, pv) -> dict:
    return {r["business_key"]: r for r in by_row(client, ed, pv)}


def live(db, number):
    db.expire_all()
    return db.scalar(select(FunctionRequirement).where(FunctionRequirement.fr_number == number))


def linked(db, f) -> set[int]:
    db.expire_all()
    return set(db.scalars(select(BrFunctionRequirement.br_id).where(BrFunctionRequirement.fr_id == f.fr_id,
                                                                  BrFunctionRequirement.is_active)))


def rate(client, ed, f, cx):
    cur = client.get(f"{API}/function-requirements/{f['fr_id']}", headers=ed["h"]).json()
    r = client.patch(f"{API}/function-requirements/{f['fr_id']}/complexity", headers=ed["h"],
                     json={"row_version": cur["row_version"], "complexity_code": cx})
    assert r.status_code == 200, r.text
    return r.json()


# ---- §10.1: an unchanged export round-trips ------------------------------------------------------------------------

def test_an_unchanged_export_round_trips_as_all_match(client, ed, world, db):
    """P2: a frozen baseline, an FR tagged outside its BRs, a rated FR and an FR holding a hidden status — all MATCH."""
    s = scope(client, ed)
    so = sol(client, ed, "Portal", "TECHNOLOGY")
    assert link(client, ed, s["br1"]["br_id"], so["solution_id"]).status_code == 201
    tagged = fr(client, ed, "Order form", "FORM", so["solution_id"])
    lnk(client, ed, s["br2"], tagged)                       # linked only to br2: tagged outside its BRs
    rate(client, ed, tagged, "MEDIUM")
    held = fr(client, ed, "Daily report", "REPORT")
    assert patch(client, ed, held, fulfilment_status_code="DEFERRED").status_code == 200
    r = client.post(f"{API}/projects/{ed['p']}/code-lists/hidden", headers=world["admin"],
                    json={"code_id": code_id(db, "FR_FULFILMENT_STATUS", "DEFERRED")})
    assert r.status_code == 201, r.text
    freeze(client, ed)
    pv = staged(client, ed, export(client, ed))
    b = pv["batch"] if "batch" in pv else pv
    assert (b["error_count"], b["insert_count"], b["update_count"]) == (0, 0, 0), by_row(client, ed, pv)
    assert {r["verdict"] for r in by_row(client, ed, pv)} == {"MATCH"}
    assert live(db, held["fr_number"]).fulfilment_status_code == "DEFERRED"


# ---- §10.4, §10.5: a new FR with links and a solution --------------------------------------------------------------

def test_a_new_row_creates_one_fr_linked_to_both_brs_and_tagged(client, ed, db):
    """§10.4 (P3a): BR-a, BR-b and a solution answering only BR-b → one FR, both links, the tag."""
    s = scope(client, ed)
    so = sol(client, ed, "Portal", "TECHNOLOGY")
    assert link(client, ed, s["br2"]["br_id"], so["solution_id"]).status_code == 201
    b1, b2 = s["br1"]["br_number"], s["br2"]["br_number"]
    body = edit(export(client, ed), lambda ws, col, rows, add: add(
        fr_name="Order form", fr_type_code="FORM", br_numbers=f"{b1}, {b2}, {b1}", solution_number=so["solution_number"],
        complexity_code="SIMPLE", description="Entry screen"))
    pv = staged(client, ed, body)
    row = by_row(client, ed, pv)[0]
    assert row["verdict"] == "INSERT" and not row["errors"], row
    r = commit(client, ed, pv, ack=1)
    assert r.status_code == 200, r.text
    f = db.scalar(select(FunctionRequirement).where(FunctionRequirement.fr_name == "Order form"))
    assert linked(db, f) == {s["br1"]["br_id"], s["br2"]["br_id"]}
    assert f.solution_id == so["solution_id"] and f.complexity_code == "SIMPLE" and f.effort_days is not None
    assert f.description == "Entry screen" and f.fulfilment_status_code == "SPECIFIED"


def test_a_solution_answering_none_of_the_rows_brs_is_an_error(client, ed):
    """§10.5."""
    s = scope(client, ed)
    so = sol(client, ed, "Portal", "TECHNOLOGY")
    assert link(client, ed, s["br2"]["br_id"], so["solution_id"]).status_code == 201
    body = edit(export(client, ed), lambda ws, col, rows, add: add(
        fr_name="Order form", fr_type_code="FORM", br_numbers=s["br1"]["br_number"],
        solution_number=so["solution_number"]))
    row = by_row(client, ed, staged(client, ed, body))[0]
    assert codes_of(row) == ["SOLUTION_ANSWERS_NO_BR"]


def test_every_row_that_validates_clean_commits(client, ed, db):
    """N1's property test: a mix of new rows and edits, every clean one written, nothing refused at commit."""
    s = scope(client, ed)
    so = sol(client, ed, "Portal", "TECHNOLOGY")
    assert link(client, ed, s["br1"]["br_id"], so["solution_id"]).status_code == 201
    a = fr(client, ed, "Invoice batch", "BATCH")
    b1 = s["br1"]["br_number"]

    def change(ws, col, rows, add):
        put(ws, rows[a["fr_number"]], col["fr_type_code"], "INTERFACE")
        put(ws, rows[a["fr_number"]], col["complexity_code"], "COMPLEX")
        put(ws, rows[a["fr_number"]], col["br_numbers"], b1)
        put(ws, rows[a["fr_number"]], col["solution_number"], so["solution_number"])
        put(ws, rows[a["fr_number"]], col["status_code"], "BUILT")
        for i, t in enumerate(("FORM", "REPORT", "BATCH", "MIGRATION")):
            add(fr_name=f"New {i}", fr_type_code=t, complexity_code="SIMPLE", br_numbers=b1,
                solution_number=so["solution_number"], status_code="TESTED")
    pv = staged(client, ed, edit(export(client, ed), change))
    rows = by_row(client, ed, pv)
    assert all(not r["errors"] for r in rows), rows
    r = commit(client, ed, pv, ack=4)
    assert r.status_code == 200, r.text
    f = live(db, a["fr_number"])
    assert (f.fr_type_code, f.complexity_code, f.solution_id, f.fulfilment_status_code) == \
        ("INTERFACE", "COMPLEX", so["solution_id"], "BUILT")
    assert db.scalar(select(FunctionRequirement).where(FunctionRequirement.fr_name == "New 3")).fr_type_code \
        == "MIGRATION"


# ---- §10.6: links are "BRs to add" ---------------------------------------------------------------------------------

def test_links_are_add_only_and_an_old_export_never_restores_a_removed_link(client, ed, db):
    s = scope(client, ed)
    b1, b2 = s["br1"], s["br2"]
    f = fr(client, ed, "Order form", "FORM")
    lnk(client, ed, b1, f)
    old = export(client, ed)

    # a row only adding BR-2 is an UPDATE whose preview says "links added"
    pv = staged(client, ed, edit(old, lambda ws, col, rows, add: put(
        ws, rows[f["fr_number"]], col["br_numbers"], f"{b1['br_number']}, {b2['br_number']}")))
    row = by_key(client, ed, pv)[f["fr_number"]]
    assert row["verdict"] == "UPDATE" and [(c["column"], c["after"]) for c in row["changes"]] == \
        [("links_added", b2["br_number"])]

    # omitting a linked BR keeps the link and warns
    pv = staged(client, ed, edit(old, lambda ws, col, rows, add: put(ws, rows[f["fr_number"]], col["br_numbers"], None)))
    row = by_key(client, ed, pv)[f["fr_number"]]
    assert row["verdict"] == "MATCH" and [w["code"] for w in row["warnings"]] == ["LINKS_ADD_ONLY"]

    # unlinked on screen since the export, and the old file still lists it → an error naming it (D1, N6)
    unlnk(client, ed, b1, next(x for x in br_rows(client, ed, b1) if x["fr_id"] == f["fr_id"]))
    row = by_key(client, ed, staged(client, ed, old))[f["fr_number"]]
    assert codes_of(row) == ["ROW_STALE"] and b1["br_number"] in row["errors"][0]["message"]
    assert linked(db, live(db, f["fr_number"])) == set()


def test_an_edited_links_at_export_cell_is_refused(client, ed):
    s = scope(client, ed)
    f = fr(client, ed, "Order form", "FORM")
    lnk(client, ed, s["br1"], f)
    body = edit(export(client, ed), lambda ws, col, rows, add: put(ws, rows[f["fr_number"]], col["br_links_at_export"],
                                                                   "| forged"))
    assert "LINKS_TOKEN_INVALID" in codes_of(by_key(client, ed, staged(client, ed, body))[f["fr_number"]])


# ---- §10.7: likely duplicates --------------------------------------------------------------------------------------

def test_a_new_row_matching_a_live_frs_name_and_type_warns(client, ed):
    scope(client, ed)
    f = fr(client, ed, "Order form", "FORM")
    body = edit(export(client, ed), lambda ws, col, rows, add: add(fr_name="  order FORM ", fr_type_code="FORM"))
    new = [r for r in by_row(client, ed, staged(client, ed, body)) if r["verdict"] == "INSERT"][0]
    assert [w["code"] for w in new["warnings"]] == ["NO_BR", "POSSIBLE_DUPLICATE"]
    assert f["fr_number"] in new["warnings"][1]["message"]


# ---- §10.10: components --------------------------------------------------------------------------------------------

def test_with_effort_off_a_complexity_change_is_an_error_naming_effort(client, ed, world):
    scope(client, ed)
    f = fr(client, ed, "Order form", "FORM")
    body = edit(export(client, ed), lambda ws, col, rows, add: put(ws, rows[f["fr_number"]], col["complexity_code"],
                                                                   "SIMPLE"))
    set_components(client, world, ed["p"], ["FUNCTION_REQUIREMENT", "TO_BE_DFD"])
    row = by_key(client, ed, staged(client, ed, body))[f["fr_number"]]
    assert codes_of(row) == ["COMPONENT_OFF"] and row["errors"][0]["params"]["component"] == "EFFORT"


# ---- §10.11, §10.12, N1, N3/N5: codes, status, blanks --------------------------------------------------------------

def test_a_hidden_code_on_a_new_row_is_an_error(client, ed, world, db):
    """§10.11 second half (an FR already holding it round-trips) is in the round-trip test above."""
    scope(client, ed)
    r = client.post(f"{API}/projects/{ed['p']}/code-lists/hidden", headers=world["admin"],
                    json={"code_id": code_id(db, "FR_FULFILMENT_STATUS", "DEFERRED")})
    assert r.status_code == 201, r.text
    body = edit(export(client, ed), lambda ws, col, rows, add: add(fr_name="Daily report", fr_type_code="REPORT",
                                                                   status_code="DEFERRED"))
    assert codes_of(by_row(client, ed, staged(client, ed, body))[0]) == ["CODE_UNKNOWN"]


def test_a_new_row_needs_a_type_and_its_typed_status_is_honoured(client, ed, db):
    """N1 and §10.12 (Q-7)."""
    scope(client, ed)
    body = edit(export(client, ed), lambda ws, col, rows, add: (add(fr_name="No type"),
                                                                  add(fr_name="Done", fr_type_code="FORM",
                                                                      status_code="TESTED")))
    rows = by_row(client, ed, staged(client, ed, body))
    assert codes_of(rows[0]) == ["REQUIRED"] and not rows[1]["errors"]
    pv = staged(client, ed, edit(export(client, ed), lambda ws, col, rows, add: add(
        fr_name="Done", fr_type_code="FORM", status_code="TESTED")))
    assert commit(client, ed, pv, ack=1).status_code == 200
    assert db.scalar(select(FunctionRequirement).where(FunctionRequirement.fr_name == "Done")).fulfilment_status_code \
        == "TESTED"


def test_a_blank_solution_or_complexity_on_an_existing_fr_is_unchanged(client, ed, db):
    """Q-5, N3, N5: never "cleared" in the preview, and nothing is written."""
    s = scope(client, ed)
    so = sol(client, ed, "Portal", "TECHNOLOGY")
    assert link(client, ed, s["br1"]["br_id"], so["solution_id"]).status_code == 201
    f = fr(client, ed, "Order form", "FORM", so["solution_id"])
    rate(client, ed, f, "MEDIUM")

    def blank(ws, col, rows, add):
        put(ws, rows[f["fr_number"]], col["solution_number"], None)
        put(ws, rows[f["fr_number"]], col["complexity_code"], None)
    row = by_key(client, ed, staged(client, ed, edit(export(client, ed), blank)))[f["fr_number"]]
    assert row["verdict"] == "MATCH" and row["changes"] == []


def test_a_changed_effort_days_cell_is_ignored_with_a_warning(client, ed):
    scope(client, ed)
    f = fr(client, ed, "Order form", "FORM")
    rate(client, ed, f, "SIMPLE")
    body = edit(export(client, ed), lambda ws, col, rows, add: put(ws, rows[f["fr_number"]], col["effort_days"], "99"))
    row = by_key(client, ed, staged(client, ed, body))[f["fr_number"]]
    assert row["verdict"] == "MATCH" and [w["code"] for w in row["warnings"]] == ["NO_BR", "EFFORT_DERIVED"]


def test_with_the_fr_component_off_template_export_and_validate_are_409(client, ed, world):
    """§10.10 (S1) for the FR sheet; the issue sheet's is in the issue tests."""
    scope(client, ed)
    body = export(client, ed)
    set_components(client, world, ed["p"], ["TO_BE_DFD"])
    assert client.get(URL.format(p=ed["p"], what="export"), headers=ed["h"]).status_code == 409
    assert client.get(f"{API}/projects/{ed['p']}/bulk/templates/function-requirements",
                      headers=ed["h"]).status_code == 409
    staged(client, ed, body, expect=409)
