"""Slice 1e-a: the BUSINESS_REQUIREMENT import target (spec docs/slice1e_imports_spec.md §7.0, §7.1; design §7.11,
R2-D1, criteria 106, 117). Grade 2: a wrong row reaches a client deliverable but is an audited, correctable update.
Every check is proven on a real export of a real chart, edited the way a consultant would edit it."""
import io

import openpyxl
from sqlalchemy import select

from models import BrRequirementLine, BusinessRequirement
from tests.test_baseline import freeze, scope
from tests.test_bulk_import import batch, codes_of, commit
from tests.test_scope_api import API, ed  # noqa: F401 — ed is a fixture

URL = API + "/projects/{p}/bulk/business-requirements/{what}"
TAG = "sheet:BUSINESS_REQUIREMENT"


def export(client, ed, p=None) -> bytes:
    r = client.get(URL.format(p=p or ed["p"], what="export"), headers=ed["h"])
    assert r.status_code == 200, r.text
    return r.content


def put(ws, row: int, column: int, value) -> None:
    """Set a cell — `ws.cell(r, c, None)` would leave it as it was (openpyxl writes only a non-None value)."""
    ws.cell(row, column).value = value


def edit(body: bytes, fn) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(body))
    ws = next(s for s in wb.worksheets if s.cell(1, 1).value == TAG)
    col = {ws.cell(1, i).value: i for i in range(2, ws.max_column + 1)}
    rows = {ws.cell(r, col["br_number"]).value: r for r in range(3, ws.max_row + 1)}
    fn(ws, col, rows)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def staged(client, ed, body):
    r = client.post(URL.format(p=ed["p"], what="validate"), content=body,
                    headers={**ed["h"], "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             "X-VB-File-Name": "brs.xlsx"})
    assert r.status_code == 201, r.text
    return r.json()


def rows_of(client, ed, pv) -> dict:
    return {r["business_key"]: r for r in batch(client, ed, pv)["rows"]}


def live(db, number):
    db.expire_all()
    return db.scalar(select(BusinessRequirement).where(BusinessRequirement.br_number == number))


def test_an_unchanged_export_round_trips_as_all_match_even_after_a_freeze(client, ed, db):
    """Spec §10.1 (P2): a frozen project's BRs are BASELINED — a status no hand may set — and still round-trip."""
    s = scope(client, ed)
    freeze(client, ed)
    pv = staged(client, ed, export(client, ed))
    b = pv["batch"] if "batch" in pv else pv
    assert (b["error_count"], b["insert_count"], b["update_count"]) == (0, 0, 0)
    assert {r["verdict"] for r in rows_of(client, ed, pv).values()} == {"MATCH"}
    assert live(db, s["br1"]["br_number"]).status_code == "BASELINED"


def test_a_changed_statement_updates_only_that_br(client, ed, db):
    """Spec §10.2: the statement changes; the step, its lines and the other BR do not."""
    s = scope(client, ed)
    n1, n2 = s["br1"]["br_number"], s["br2"]["br_number"]
    lines_before = db.scalar(select(BrRequirementLine.br_requirement_line_id).limit(1))
    body = edit(export(client, ed), lambda ws, col, rows: put(ws, rows[n1], col["br_statement"], "Take every order"))
    pv = staged(client, ed, body)
    rows = rows_of(client, ed, pv)
    assert rows[n1]["verdict"] == "UPDATE" and rows[n2]["verdict"] == "MATCH"
    assert [c["column"] for c in rows[n1]["changes"]] == ["br_statement"]
    r = commit(client, ed, pv)
    assert r.status_code == 200, r.text
    assert live(db, n1).br_statement == "Take every order"
    assert live(db, n2).br_statement == s["br2"].get("br_statement")
    assert db.scalar(select(BrRequirementLine.br_requirement_line_id).limit(1)) == lines_before


def test_a_blank_or_foreign_br_number_is_an_error_and_nothing_is_written(client, ed, db):
    """Criterion 117b, spec §10.3: every BR exists with its step; an import never creates one."""
    s = scope(client, ed)
    n1 = s["br1"]["br_number"]

    def change(ws, col, rows):
        put(ws, rows[n1], col["br_number"], None)
        put(ws, rows[s["br2"]["br_number"]], col["br_number"], "BR-9999")
    rows = batch(client, ed, staged(client, ed, edit(export(client, ed), change)))["rows"]
    assert sorted(c for r in rows for c in codes_of(r)) == ["BR_NUMBER_REQUIRED", "NOT_FOUND"]
    assert db.scalar(select(BusinessRequirement).where(BusinessRequirement.br_number == "BR-9999")) is None


def test_a_br_moved_since_export_is_an_error(client, ed, db):
    """Criterion 117 (R2-D1): hier_code is a cross-check only; a different one is "moved since export"."""
    s = scope(client, ed)
    n1 = s["br1"]["br_number"]

    def change(ws, col, rows):
        put(ws, rows[n1], col["hier_code"], "09.09.09")
        put(ws, rows[n1], col["br_statement"], "Changed")
    row = rows_of(client, ed, staged(client, ed, edit(export(client, ed), change)))[n1]
    assert codes_of(row) == ["MOVED_SINCE_EXPORT"]
    assert live(db, n1).br_statement != "Changed"


def test_a_freeze_since_export_is_named_as_the_reason(client, ed):
    """Spec §7.0 (D6): the stale row says a baseline was frozen, not just "changed"."""
    s = scope(client, ed)
    body = export(client, ed)
    freeze(client, ed)
    row = rows_of(client, ed, staged(client, ed, body))[s["br1"]["br_number"]]
    assert codes_of(row) == ["FROZEN_SINCE_EXPORT"]


def test_status_and_statement_refusals_run_only_on_a_changed_value(client, ed):
    """Spec §7.0, §7.1: BASELINED by hand, an unknown status and a cleared statement are errors; a renamed step is a
    warning and nothing else."""
    s = scope(client, ed)
    n1, n2 = s["br1"]["br_number"], s["br2"]["br_number"]
    client.patch(f"{API}/business-requirements/{s['br2']['br_id']}", headers=ed["h"],
                 json={"row_version": s["br2"]["row_version"], "br_statement": "Bill it"})

    def change(ws, col, rows):
        put(ws, rows[n1], col["status_code"], "BASELINED")
        put(ws, rows[n2], col["br_statement"], None)
        put(ws, rows[n2], col["step_name"], "A new name")
    body = edit(export(client, ed), change)
    rows = rows_of(client, ed, staged(client, ed, body))
    assert codes_of(rows[n1]) == ["SYSTEM_SET"]
    assert codes_of(rows[n2]) == ["REQUIRED"]
    assert [w["code"] for w in rows[n2]["warnings"]] == ["STEP_NAME_IGNORED"]

    rows = rows_of(client, ed, staged(client, ed, edit(export(client, ed),
                                                       lambda ws, col, r: put(ws, r[n1], col["status_code"], "NOPE"))))
    assert codes_of(rows[n1]) == ["CODE_UNKNOWN"]


def test_a_row_without_its_token_must_be_exported_first(client, ed):
    s = scope(client, ed)
    n1 = s["br1"]["br_number"]
    row = rows_of(client, ed, staged(client, ed, edit(export(client, ed),
                                                      lambda ws, col, r: put(ws, r[n1], col["row_token"], None))))[n1]
    assert codes_of(row) == ["EXPORT_FIRST"]
