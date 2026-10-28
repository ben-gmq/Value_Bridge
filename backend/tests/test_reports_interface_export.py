"""Interface list to Excel (docs/slice1f_spec.md §3.5, function 61). Grade 2: the file goes to a vendor as a scope
document; a short filtered list or a formula in a cell misleads them, and is corrected once someone checks. Proven on
one built fixture: the stamp, one row per interface with what it lacks, the filtered-out count (D6), a formula-lead
name kept as text, the effort column blank with EFFORT off (A1), and 404 for an outsider.

Deliberately uncovered (spec §5): column widths and look in Excel itself; very large lists."""
import io

import openpyxl
from sqlalchemy import text

from tests.test_br_fr import set_components
from tests.test_effort_api import rate
from tests.test_function_requirement_api import fr
from tests.test_interface_api import put, world_
from tests.test_scope_api import API, ed  # noqa: F401 — ed is a fixture


def export(client, ed, headers=None, expect=200, **q):
    r = client.get(f"{API}/projects/{ed['p']}/interfaces/export", headers=headers or ed["rv"], params=q)
    assert r.status_code == expect, r.text
    if expect != 200:
        return None
    assert r.headers["content-disposition"] == f'attachment; filename="vb-interfaces-{ed["p"]}.xlsx"'
    return openpyxl.load_workbook(io.BytesIO(r.content)).active


def table(ws):
    """The header row and the data rows under it, as lists of values."""
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    start = next(i for i, r in enumerate(rows) if r[0] == "Interface")
    return rows[start], rows[start + 1:]


def test_stamp_rows_and_what_an_incomplete_interface_lacks(client, ed):
    erp, crm, order, _, f = world_(client, ed)
    put(client, ed, f, erp, crm, [{"data_entity_id": order["data_entity_id"], "note": "header and lines"}],
        pattern_code="PUSH", method_code="API", frequency_code="DAILY")
    bare = fr(client, ed, "=HYPERLINK(\"x\")", "INTERFACE")                      # never described, formula lead
    assert rate(client, ed, f, "COMPLEX").status_code == 200
    ws = export(client, ed)
    assert ws["A1"].font.bold and ws["A2"].value == "LIVE — not a baseline"
    assert ws["A3"].value.startswith("Generated ") and " ICT by " in ws["A3"].value
    assert ws["A4"].value == "Filters: none"
    header, rows = table(ws)
    assert header[0:4] == ["Interface", "Name", "Type", "Source → target"] and len(rows) == 2
    done = next(r for r in rows if r[0] == f["fr_number"])
    assert done[3:8] == ["SAP-ECC → SF-CRM", "Push", "API", "Daily", f"{order['de_number']} Sales order (header and lines)"]
    assert done[11] > 0 and done[12] == "Complete"
    todo = next(r for r in rows if r[0] == bare["fr_number"])
    assert todo[1] == '=HYPERLINK("x")' and todo[3] is None
    assert todo[12] == "Missing: endpoints, method, data entities"
    cell = next(c for c in ws["B"] if c.value == '=HYPERLINK("x")')
    assert cell.data_type == "s" and cell.quotePrefix                            # text, never a formula


def test_a_pair_filter_prints_itself_and_the_interfaces_it_cannot_match(client, ed):
    """D6: a filtered file says what it was filtered by and how many undescribed interfaces it left out."""
    erp, crm, order, _, f = world_(client, ed)
    put(client, ed, f, erp, crm, [order], method_code="API")
    fr(client, ed, "Not yet described", "INTERFACE")
    ws = export(client, ed, source=erp["application_id"], status="SPECIFIED")
    assert ws["A4"].value == "Filters: source SAP-ECC, target any, status SPECIFIED"
    assert ws["A5"].value == "1 interfaces not yet described are not in this filtered list"
    _, rows = table(ws)
    assert [r[0] for r in rows] == [f["fr_number"]]


def test_effort_off_leaves_the_column_blank_and_an_outsider_gets_404(client, ed, world, db):
    """EFFORT cannot go off while an FR is rated, so a leftover figure is planted to prove the column ignores it."""
    erp, crm, order, _, f = world_(client, ed)
    put(client, ed, f, erp, crm, [order])
    set_components(client, world, ed["p"], ["FUNCTION_REQUIREMENT", "APPLICATION", "SOLUTION"])
    db.execute(text("UPDATE function_requirement SET effort_days = 5, complexity_code_id = (SELECT code_id FROM "
                    "code_master WHERE category = 'FR_COMPLEXITY' AND code = 'COMPLEX' AND project_id IS NULL) "
                    "WHERE fr_id = :f"), {"f": f["fr_id"]})
    db.commit()
    header, rows = table(export(client, ed))
    assert header[11] == "Effort (days)" and rows[0][11] is None
    export(client, ed, headers=ed["out"], expect=404)
