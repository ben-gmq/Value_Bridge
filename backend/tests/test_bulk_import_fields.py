"""Slice 4a-2 — Data Fields through the one import pipeline, and the one-workbook data model
(docs/slice4a_spec.md §3, §4 `validate data_field`, §10 criteria 9, 10, 17, 19, 20; rules 4a-R1,
R2, R10, R12…R17, R19). Grade 1 for the commit path: each rule is proven against the database.

A data-model workbook runs as two batches (S4-12): the entities sheet through `data-entities`,
then the fields sheet through `data-fields`, posting the same bytes both times."""
from contextlib import contextmanager
import io
from urllib.parse import quote

import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy import select, text
from sqlalchemy.exc import OperationalError

from database import SessionLocal
from models import DataEntity, DataField, ImportRow
from services import bulk, consistency_data, xlsx
from services import data_entity as de_service
from tests.test_bulk_import import CODES as DE_CODES, codes_of, commit, live
from tests.test_erd_api import field, pk
from tests.test_scope_api import API, ed, entity  # noqa: F401 — ed is a fixture

DF_CODES = bulk.DATA_FIELD.codes
DE_TAG, DF_TAG = "sheet:DATA_ENTITY", "sheet:DATA_FIELD"
BULK = API + "/projects/{p}/bulk/{what}"


# ---- workbooks ----------------------------------------------------------------------------

def _sheet(wb, title, tag, codes, rows):
    ws = wb.create_sheet(title)
    ws.cell(1, 1).value = tag
    for i, c in enumerate(codes, start=2):
        ws.cell(1, i).value = c
        ws.cell(2, i).value = c
    for r, row in enumerate(rows, start=3):
        for i, c in enumerate(codes, start=2):
            if row.get(c) is not None:
                ws.cell(r, i).value = row[c]


def model(entities=None, fields=()) -> bytes:
    """A data-model workbook: an entities sheet (when `entities` is not None) and a fields sheet."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    if entities is not None:
        _sheet(wb, "Data entities", DE_TAG, DE_CODES, [{"de_name": n} for n in entities])
    _sheet(wb, "Data fields", DF_TAG, DF_CODES, list(fields))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def edit_fields(body: bytes, fn) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(body))
    ws = next(s for s in wb.worksheets if s.cell(1, 1).value == DF_TAG)
    col = {ws.cell(1, i).value: i for i in range(2, ws.max_column + 1)}
    fn(ws, col)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def setc(ws, col, sheet_row, code, value):
    ws.cell(sheet_row, col[code]).value = value          # ws.cell(r, c, None) would not clear it


# ---- the API ------------------------------------------------------------------------------

def post(client, ed, what, body, p=None, name="model.xlsx"):
    return client.post(BULK.format(p=p or ed["p"], what=what), content=body,
                       headers={**ed["h"], "Content-Type": xlsx.XLSX_MEDIA_TYPE,
                                "X-VB-File-Name": quote(name)})


def stage(client, ed, target, body, **kw):
    r = post(client, ed, f"{target}/validate", body, **kw)
    assert r.status_code == 201, r.text
    return r.json()


def get(client, ed, what, headers=None):
    r = client.get(BULK.format(p=ed["p"], what=what), headers=headers or ed["h"])
    assert r.status_code == 200, r.text
    return r.content


def preview(client, ed, pv):
    return client.get(f"{API}/bulk/batches/{pv['batch']['import_batch_id']}/preview", headers=ed["h"]).json()


def ok_commit(client, ed, pv):
    r = commit(client, ed, pv, ack=pv["batch"]["insert_count"] or None)
    assert r.status_code == 200, r.text
    return r.json()


def fields_of(db, p) -> dict[str, DataField]:
    """Live fields by 'DE-nnnn/name'."""
    db.expire_all()
    rows = db.execute(select(DataField, DataEntity.de_number)
                      .join(DataEntity, DataEntity.data_entity_id == DataField.data_entity_id)
                      .where(DataField.project_id == p, DataField.is_active)).all()
    return {f"{n}/{f.field_name}": f for f, n in rows}


def by_row(pv) -> dict[int, dict]:
    return {r["sheet_row_no"]: r for r in pv["rows"]}


# ---- criterion 9 (S4-12, 4a-R12, 4a-R19): a whole model in one workbook ---------------------

WHOLE = [
    # Customer: a two-part key, named by name.
    {"de_name": "Customer", "field_name": "cust_id", "data_type": "STRING", "length": 10, "pk_position": 1},
    {"de_name": "customer ", "field_name": "region", "data_type": "STRING", "pk_position": 2},
    # Order: named by number on one row and by name on another — one parent (4a-R12).
    {"de_number": "DE-0002", "field_name": "order_id", "data_type": "INTEGER", "pk_position": 1},
    # A composite FK: label a, its target named by ref_de_name (4a-R19).
    {"de_name": "Order", "field_name": "cust_id", "ref_de_name": "Customer", "ref_field_name": "cust_id",
     "fk_group": "a", "mandatory": "Y"},
    {"de_name": "Order", "field_name": "cust_region", "ref_de_name": "CUSTOMER", "ref_field_name": "Region",
     "fk_group": "a"},
    # A self-reference.
    {"de_name": "Employee", "field_name": "emp_id", "pk_position": 1},
    {"de_name": "Employee", "field_name": "manager_id", "ref_de_name": "Employee", "ref_field_name": "emp_id"},
    # The same label a under another parent: a second relationship, not the first one's.
    {"de_name": "Invoice", "field_name": "billed_cust", "ref_de_number": "DE-0001", "ref_field_name": "cust_id",
     "fk_group": "a"},
    {"de_name": "Invoice", "field_name": "billed_region", "ref_de_number": "DE-0001",
     "ref_field_name": "region", "fk_group": "a"},
]


def test_criterion_9_a_whole_model_in_one_workbook_commits_in_two_steps(client, ed, db):
    body = model(["Customer", "Order", "Employee", "Invoice"], WHOLE)

    step1 = stage(client, ed, "data-entities", body)
    assert [w["code"] for w in step1["batch"]["file_warnings"]] == ["SHEET_NOT_READ"]   # the fields sheet
    ok_commit(client, ed, step1)
    assert sorted(live(db, ed["p"])) == ["DE-0001", "DE-0002", "DE-0003", "DE-0004"]

    step2 = stage(client, ed, "data-fields", body)
    b = step2["batch"]
    assert (b["error_count"], b["insert_count"], b["sheet_name"]) == (0, 9, "Data fields"), step2["rows"]
    assert [w["code"] for w in b["file_warnings"]] == ["SHEET_NOT_READ"]               # 4a-R17
    rows = by_row(step2)
    assert rows[4]["business_key"] == "DE-0001/region"
    # 4a-R13: a row named by name shows the number and name it resolved to.
    assert rows[3]["resolved"]["parent"] == {"de_number": "DE-0001", "de_name": "Customer", "by_name": True}
    assert rows[5]["resolved"]["parent"]["by_name"] is False
    # 4a-R1: the preview names the relationship each new FK row joins.
    assert rows[6]["resolved"]["relationship"] == {"kind": "new", "label": "a", "to": "DE-0001",
                                                   "with": ["cust_region"]}
    assert rows[6]["resolved"]["ref"] == {"de_number": "DE-0001", "de_name": "Customer", "by_name": True}
    ok_commit(client, ed, step2)

    f = fields_of(db, ed["p"])
    assert len(f) == 9
    cust, region = f["DE-0001/cust_id"], f["DE-0001/region"]
    assert (cust.pk_ordinal, region.pk_ordinal, cust.length_val) == (1, 2, 10)
    o1, o2 = f["DE-0002/cust_id"], f["DE-0002/cust_region"]
    assert (o1.ref_data_field_id, o2.ref_data_field_id) == (cust.data_field_id, region.data_field_id)
    assert o1.fk_group_no == o2.fk_group_no == 1 and o1.is_mandatory is True     # one relationship per label
    me = f["DE-0003/manager_id"]
    assert (me.ref_data_entity_id, me.ref_data_field_id) == (f["DE-0003/emp_id"].data_entity_id,
                                                             f["DE-0003/emp_id"].data_field_id)
    i1, i2 = f["DE-0004/billed_cust"], f["DE-0004/billed_region"]
    assert i1.fk_group_no == i2.fk_group_no == 1 and i1.data_entity_id != o1.data_entity_id
    assert [x.seq_no for x in (o1, o2)] == [2, 3]                                  # sheet order
    erd = client.get(f"{API}/projects/{ed['p']}/erd", headers=ed["h"]).json()
    rels = sorted((r["child_data_entity_id"], r["parent_data_entity_id"], r["fk_group_no"], r["label"])
                  for r in erd["relationships"])
    assert [x[3] for x in rels] == ["cust_id, cust_region", "manager_id", "billed_cust, billed_region"]


def test_a_number_and_a_name_that_disagree_or_find_nothing_are_row_errors(client, ed, three):
    pv = stage(client, ed, "data-fields", model(None, [
        {"de_number": "DE-0001", "de_name": "Order", "field_name": "x"},
        {"de_number": "DE-0009", "field_name": "x"},
        {"de_name": "Nobody", "field_name": "x"},
        {"de_name": "Order", "field_name": "y", "ref_de_number": "DE-0003", "ref_de_name": "Customer"},
        {"field_name": "orphan"},
        {"de_name": "Order", "field_name": "z", "ref_field_name": "id"},
    ]))
    rows = by_row(pv)
    assert codes_of(rows[3]) == ["ENTITY_MISMATCH"] and rows[3]["errors"][0]["column"] == "de_name"
    assert codes_of(rows[4]) == ["ENTITY_NOT_FOUND"]
    assert codes_of(rows[5]) == ["ENTITY_NOT_FOUND"]
    assert codes_of(rows[6]) == ["ENTITY_MISMATCH"] and rows[6]["errors"][0]["column"] == "ref_de_name"
    assert codes_of(rows[7]) == ["REQUIRED"]
    assert codes_of(rows[8]) == ["REF_NEEDS_ENTITY"]


@pytest.fixture
def three(client, ed):
    return [entity(client, ed, n) for n in ("Customer", "Order", "Product")]


def test_4a_r12_keys_use_the_resolved_parent_never_the_raw_cell(client, ed, three):
    """One parent by number and by name is one parent: the same field twice is a duplicate."""
    pv = stage(client, ed, "data-fields", model(None, [
        {"de_number": "DE-0002", "field_name": "Order ID"},
        {"de_name": " order", "field_name": "order id "},
    ]))
    rows = pv["rows"]
    assert rows[0]["business_key"] == "DE-0002/order id" and rows[0]["errors"] == []
    assert codes_of(rows[1]) == ["DUPLICATE_KEY"] and rows[1]["errors"][0]["params"]["first_row"] == 3


# ---- criterion 10: swaps are validation errors, never a mid-commit crash -------------------

def test_criterion_10_a_swap_of_two_field_names_or_two_pk_positions_is_a_validation_error(client, ed, db, three):
    de = three[0]
    pk(client, ed, de, "a", 1)
    pk(client, ed, de, "b", 2)
    exported = get(client, ed, "data-fields/export")

    def swap_names(ws, col):
        setc(ws, col, 3, "field_name", "b")
        setc(ws, col, 4, "field_name", "a")

    def swap_pks(ws, col):
        setc(ws, col, 3, "pk_position", 2)
        setc(ws, col, 4, "pk_position", 1)
    for change, code in ((swap_names, "FIELD_MISMATCH"), (swap_pks, "PK_TAKEN")):
        pv = stage(client, ed, "data-fields", edit_fields(exported, change))
        assert [codes_of(r) for r in pv["rows"]] == [[code], [code]]
        assert commit(client, ed, pv).status_code == 422
    f = fields_of(db, ed["p"])
    assert (f["DE-0001/a"].pk_ordinal, f["DE-0001/b"].pk_ordinal) == (1, 2)


def test_4a_r2_a_new_pk_position_a_live_field_holds_is_an_error(client, ed, three):
    pk(client, ed, three[0], "id", 1)
    pv = stage(client, ed, "data-fields", model(None, [
        {"de_number": "DE-0001", "field_name": "code", "pk_position": 1},
        {"de_number": "DE-0001", "field_name": "x", "pk_position": 2},
        {"de_number": "DE-0001", "field_name": "y", "pk_position": 2},
        {"de_number": "DE-0001", "field_name": "z", "pk_position": 3, "mandatory": "N"},
    ]))
    assert [codes_of(r) for r in pv["rows"]] == [["PK_TAKEN"], [], ["DUPLICATE_PK"], ["PK_OPTIONAL"]]


# ---- criterion 20 (S4-12, 4a-R10): a fields step that fails after its entities committed --

def test_criterion_20_entities_stay_live_and_the_corrected_file_goes_to_the_fields_step_alone(client, ed, db):
    broken = model(["Customer", "Order"], [
        {"de_name": "Customer", "field_name": "id", "data_type": "NOT_A_TYPE"},
        {"de_name": "Order", "field_name": "id"}])
    ok_commit(client, ed, stage(client, ed, "data-entities", broken))
    failed = stage(client, ed, "data-fields", broken)
    assert failed["batch"]["error_count"] == 1 and codes_of(failed["rows"][0]) == ["CODE_UNKNOWN"]
    assert commit(client, ed, failed).status_code == 422

    assert fields_of(db, ed["p"]) == {} and sorted(live(db, ed["p"])) == ["DE-0001", "DE-0002"]
    listed = client.get(f"{API}/projects/{ed['p']}/data-entities", headers=ed["h"]).json()
    assert [(e["de_number"], e["live_field_count"]) for e in listed] == [("DE-0001", 0), ("DE-0002", 0)]
    assert [e["de_number"] for e in consistency_data.entities_without_fields(db, ed["p"])] == \
        ["DE-0001", "DE-0002"]

    corrected = model(["Customer", "Order"], [
        {"de_name": "Customer", "field_name": "id", "data_type": "STRING"},
        {"de_name": "Order", "field_name": "id"}])
    pv = stage(client, ed, "data-fields", corrected)
    assert pv["batch"]["error_count"] == 0                     # no "name already exists": not re-read
    assert [w["code"] for w in pv["batch"]["file_warnings"]] == ["SHEET_NOT_READ"]
    ok_commit(client, ed, pv)
    assert sorted(fields_of(db, ed["p"])) == ["DE-0001/id", "DE-0002/id"]
    assert consistency_data.entities_without_fields(db, ed["p"]) == []
    listed = client.get(f"{API}/projects/{ed['p']}/data-entities", headers=ed["h"]).json()
    assert [e["live_field_count"] for e in listed] == [1, 1]


def test_entities_without_fields_lists_live_entities_with_no_live_field(client, ed, db, three):
    field(client, ed, three[1], "id")
    gone = field(client, ed, three[2], "old")
    client.delete(f"{API}/data-fields/{gone['data_field_id']}", headers=ed["h"],
                  params={"row_version": gone["row_version"]})
    assert consistency_data.entities_without_fields(db, ed["p"]) == [
        {"data_entity_id": three[0]["data_entity_id"], "de_number": "DE-0001", "de_name": "Customer"},
        {"data_entity_id": three[2]["data_entity_id"], "de_number": "DE-0003", "de_name": "Product"}]


# ---- criterion 5, for fields: a token whose field was renamed, moved or retired ------------

def test_a_token_whose_field_was_renamed_or_retired_is_an_error_never_an_insert(client, ed, db, three):
    a = field(client, ed, three[0], "a")
    b = field(client, ed, three[0], "b")
    exported = get(client, ed, "data-fields/export")
    client.patch(f"{API}/data-fields/{a['data_field_id']}", headers=ed["h"],
                 json={"row_version": a["row_version"], "field_name": "a2"})
    assert client.delete(f"{API}/data-fields/{b['data_field_id']}", headers=ed["h"],
                         params={"row_version": b["row_version"]}).status_code == 204
    pv = stage(client, ed, "data-fields", exported)
    assert [codes_of(r) for r in pv["rows"]] == [["FIELD_MISMATCH"], ["ROW_RETIRED"]]
    assert pv["batch"]["insert_count"] == 0


def test_a_token_row_moved_to_another_entity_is_an_error(client, ed, three):
    field(client, ed, three[0], "a")
    pv = stage(client, ed, "data-fields", edit_fields(get(client, ed, "data-fields/export"),
                                                      lambda ws, col: (setc(ws, col, 3, "de_number", "DE-0002"),
                                                                       setc(ws, col, 3, "de_name", None))))
    assert codes_of(pv["rows"][0]) == ["FIELD_MISMATCH"]


def test_a_new_row_for_an_existing_field_is_export_first(client, ed, three):
    field(client, ed, three[0], "Code")
    pv = stage(client, ed, "data-fields", model(None, [{"de_number": "DE-0001", "field_name": " code"}]))
    assert codes_of(pv["rows"][0]) == ["EXPORT_FIRST"]


def test_a_field_row_going_stale_between_validate_and_commit_applies_nothing(client, ed, db, three):
    a = field(client, ed, three[0], "a")
    pv = stage(client, ed, "data-fields", edit_fields(
        get(client, ed, "data-fields/export"),
        lambda ws, col: (setc(ws, col, 3, "description", "mine"), ws.cell(4, col["de_number"], "DE-0001"),
                         ws.cell(4, col["field_name"], "new one"))))
    assert [r["verdict"] for r in pv["rows"]] == ["UPDATE", "INSERT"]
    client.patch(f"{API}/data-fields/{a['data_field_id']}", headers=ed["h"],
                 json={"row_version": a["row_version"], "description": "theirs"})
    r = commit(client, ed, pv, ack=1)
    assert r.status_code == 422
    assert codes_of(preview(client, ed, pv)["rows"][0]) == ["ROW_STALE"]
    f = fields_of(db, ed["p"])
    assert sorted(f) == ["DE-0001/a"] and f["DE-0001/a"].description == "theirs"


# ---- criterion 8 and 17, for fields: a round trip; before-values ---------------------------

def test_export_then_reimport_unchanged_is_all_match_and_bumps_nothing(client, ed, db, three):
    cid = pk(client, ed, three[0], "id", 1, data_type_code="DECIMAL", precision_val=10, scale_val=2,
             description="-1 minus first")
    field(client, ed, three[1], "cust", is_mandatory=False, ref_data_entity_id=three[0]["data_entity_id"],
          ref_data_field_id=cid["data_field_id"])
    before = {k: f.row_version for k, f in fields_of(db, ed["p"]).items()}
    pv = stage(client, ed, "data-fields", get(client, ed, "data-fields/export"))
    assert [r["verdict"] for r in pv["rows"]] == ["MATCH", "MATCH"], pv["rows"]
    ok_commit(client, ed, pv)
    after = fields_of(db, ed["p"])
    assert {k: f.row_version for k, f in after.items()} == before
    assert after["DE-0001/id"].description == "-1 minus first"


def test_commit_stores_before_values_for_each_update_row(client, ed, db, three):
    a = field(client, ed, three[0], "a", description="old", data_type_code="STRING", length_val=5)
    pv = stage(client, ed, "data-fields", edit_fields(
        get(client, ed, "data-fields/export"),
        lambda ws, col: (setc(ws, col, 3, "description", None), setc(ws, col, 3, "length", 8),
                         setc(ws, col, 3, "mandatory", "Y"))))
    [row] = pv["rows"]
    assert {c["column"]: (c["before"], c["after"], c["cleared"]) for c in row["changes"]} == {
        "length": (5, 8, False), "mandatory": (None, True, False), "description": ("old", None, True)}
    ok_commit(client, ed, pv)
    stored = db.scalars(select(ImportRow).where(
        ImportRow.import_batch_id == pv["batch"]["import_batch_id"])).one()
    assert stored.payload["before"]["description"] == "old" and stored.payload["before"]["length"] == 5
    assert stored.payload["before"]["row_version"] == a["row_version"]
    f = fields_of(db, ed["p"])["DE-0001/a"]
    assert (f.description, f.length_val, f.is_mandatory) == (None, 8, True)


# ---- 4a-R1: a numeric fk_group is a stored relationship; anything else is a label ----------

def test_4a_r1_numeric_groups_must_exist_and_join_labels_start_new_ones(client, ed, db, three):
    cust, order = three[0], three[1]
    k1 = pk(client, ed, cust, "id", 1)
    k2 = pk(client, ed, cust, "region", 2)
    field(client, ed, order, "cust_id", ref_data_entity_id=cust["data_entity_id"],
          ref_data_field_id=k1["data_field_id"])                                  # relationship 1

    pv = stage(client, ed, "data-fields", model(None, [
        {"de_number": "DE-0002", "field_name": "bad", "ref_de_number": "DE-0001", "fk_group": "7"}]))
    assert codes_of(pv["rows"][0]) == ["FK_GROUP_NOT_FOUND"]

    def add(ws, col):
        for i, (name, ref_field, group) in enumerate((("cust_region", "region", 1),       # joins rel 1
                                                       ("alt_id", "id", "b"), ("alt_region", "region", "b"))):
            for k, v in {"de_number": "DE-0002", "field_name": name, "ref_de_number": "DE-0001",
                         "ref_field_name": ref_field, "fk_group": group}.items():
                ws.cell(6 + i, col[k]).value = v
    pv = stage(client, ed, "data-fields", edit_fields(get(client, ed, "data-fields/export"), add))
    rows = by_row(pv)
    assert pv["batch"]["error_count"] == 0, pv["rows"]
    assert rows[6]["resolved"]["relationship"] == {"kind": "existing", "group": 1, "to": "DE-0001",
                                                   "with": ["cust_id"]}
    assert rows[7]["resolved"]["relationship"] == {"kind": "new", "label": "b", "to": "DE-0001",
                                                   "with": ["alt_region"]}
    ok_commit(client, ed, pv)
    f = fields_of(db, ed["p"])
    assert f["DE-0002/cust_region"].fk_group_no == f["DE-0002/cust_id"].fk_group_no == 1
    assert f["DE-0002/alt_id"].fk_group_no == f["DE-0002/alt_region"].fk_group_no == 2
    assert f["DE-0002/cust_region"].ref_data_field_id == k2["data_field_id"]


def test_one_relationship_cannot_point_at_one_field_twice(client, ed, three):
    pk(client, ed, three[0], "id", 1)
    pv = stage(client, ed, "data-fields", model(None, [
        {"de_number": "DE-0002", "field_name": f, "ref_de_number": "DE-0001", "ref_field_name": "id",
         "fk_group": "a"} for f in ("x", "y")]))
    assert [codes_of(r) for r in pv["rows"]] == [[], ["FK_FIELD_TWICE"]]


def test_a_blank_reference_removes_the_foreign_key(client, ed, db, three):
    k = pk(client, ed, three[0], "id", 1)
    field(client, ed, three[1], "cust", ref_data_entity_id=three[0]["data_entity_id"],
          ref_data_field_id=k["data_field_id"])
    pv = stage(client, ed, "data-fields", edit_fields(
        get(client, ed, "data-fields/export"),
        lambda ws, col: [setc(ws, col, 4, c, None) for c in ("ref_de_number", "ref_de_name",
                                                            "ref_field_name", "fk_group")]))
    assert [r["verdict"] for r in pv["rows"]] == ["MATCH", "UPDATE"]
    ok_commit(client, ed, pv)
    f = fields_of(db, ed["p"])["DE-0002/cust"]
    assert (f.is_foreign_key, f.ref_data_entity_id, f.fk_group_no) == (False, None, None)


# ---- 4a-R13: the resolved parent is pinned at staging --------------------------------------

def test_4a_r13_a_rename_between_validate_and_commit_is_an_error_never_a_write_elsewhere(client, ed, db):
    a, b = entity(client, ed, "Customer"), entity(client, ed, "Client")
    pv = stage(client, ed, "data-fields", model(None, [{"de_name": "Customer", "field_name": "id"}]))
    assert pv["rows"][0]["business_key"] == "DE-0001/id"
    for de, name in ((a, "Old customer"), (b, "Customer")):
        r = client.patch(f"{API}/data-entities/{de['data_entity_id']}", headers=ed["h"],
                         json={"row_version": de["row_version"], "de_name": name})
        assert r.status_code == 200
    r = commit(client, ed, pv, ack=1)
    assert r.status_code == 422
    after = preview(client, ed, pv)
    assert after["batch"]["status"] == "REJECTED"
    assert codes_of(after["rows"][0]) == ["RESOLVED_CHANGED"]
    assert after["rows"][0]["errors"][0]["params"] == {"was": "DE-0001", "now": "DE-0002"}
    assert fields_of(db, ed["p"]) == {}


def test_a_parent_retired_after_validate_rejects_the_batch(client, ed, db, three):
    pv = stage(client, ed, "data-fields", model(None, [{"de_number": "DE-0003", "field_name": "id"}]))
    client.delete(f"{API}/data-entities/{three[2]['data_entity_id']}", headers=ed["h"],
                  params={"row_version": three[2]["row_version"]})
    assert commit(client, ed, pv, ack=1).status_code == 422
    assert codes_of(preview(client, ed, pv)["rows"][0]) == ["ENTITY_NOT_FOUND"]
    assert fields_of(db, ed["p"]) == {}


# ---- 4a-R14: names are bound to the batch's project ----------------------------------------

def test_4a_r14_a_name_live_only_in_another_project_is_not_found_and_never_leaks(client, ed, world):
    other = {"h": world["admin"], "p": world["p_beta"]}
    for n in ("Ledger", "Journal", "Customer"):
        entity(client, other, n)                                   # Customer is DE-0003 over there
    r = post(client, ed, "data-fields/validate", model(None, [
        {"de_name": "Customer", "field_name": "id"},
        {"de_name": "Order", "field_name": "x", "ref_de_name": "Customer "}]))
    assert r.status_code == 201
    assert [codes_of(x) for x in r.json()["rows"]] == [["ENTITY_NOT_FOUND"], ["ENTITY_NOT_FOUND"] * 2]
    assert "DE-000" not in r.text


def test_4a_r15_a_name_that_differs_only_in_spaces_is_named_in_the_error(client, ed, three):
    entity(client, ed, "Sales Order")
    pv = stage(client, ed, "data-fields", model(None, [{"de_name": "Sales  order x", "field_name": "a"},
                                                       {"de_name": "SalesOrder", "field_name": "a"}]))
    assert pv["rows"][0]["errors"][0]["params"]["near"] is None
    assert pv["rows"][1]["errors"][0]["params"]["near"] == "DE-0004 Sales Order"


# ---- the data-model template and export (S4-12) --------------------------------------------

def test_the_data_model_template_has_both_sheets_and_one_lookup_sheet(client, ed, db):
    from tests.test_bulk_import import _with_types
    _with_types(db, ed["p"], [f"LONG_TYPE_CODE_NUMBER_{i:02d}" for i in range(14)])
    wb = openpyxl.load_workbook(io.BytesIO(get(client, ed, "templates/data-model")))
    assert [ws.cell(1, 1).value for ws in wb.worksheets[:2]] == [DE_TAG, DF_TAG]
    assert [ws.title for ws in wb.worksheets[2:]] == [xlsx.LOOKUP_SHEET]          # one, shared
    assert wb.worksheets[2].sheet_state == "hidden"
    fields_ws = wb.worksheets[1]
    assert [fields_ws.cell(1, i).value for i in range(2, 2 + len(DF_CODES))] == list(DF_CODES)
    lists = {str(dv.sqref)[0]: dv.formula1 for dv in fields_ws.data_validations.dataValidation}
    assert lists["I"] == '"Y,N"' and lists["E"].startswith(f"'{xlsx.LOOKUP_SHEET}'!")
    r = post(client, ed, "data-model/validate", get(client, ed, "templates/data-model"))
    assert r.status_code == 404                                   # validate stays per target


def test_the_data_model_export_round_trips_through_both_steps_as_all_match(client, ed, db, three):
    k = pk(client, ed, three[0], "id", 1, data_type_code="STRING", length_val=12)
    field(client, ed, three[1], "cust", ref_data_entity_id=three[0]["data_entity_id"],
          ref_data_field_id=k["data_field_id"], is_mandatory=True)
    field(client, ed, three[1], "parent_order", ref_data_entity_id=three[1]["data_entity_id"])
    exported = get(client, ed, "data-model/export", headers=ed["rv"])           # a REVIEWER may export
    wb = openpyxl.load_workbook(io.BytesIO(exported))
    assert [ws.cell(1, 1).value for ws in wb.worksheets] == [DE_TAG, DF_TAG]

    entities = stage(client, ed, "data-entities", exported)
    assert [r["verdict"] for r in entities["rows"]] == ["MATCH"] * 3
    ok_commit(client, ed, entities)
    fields = stage(client, ed, "data-fields", exported)
    assert [r["verdict"] for r in fields["rows"]] == ["MATCH"] * 3, fields["rows"]
    assert fields["batch"]["error_count"] == 0


def test_the_check_route_runs_the_file_checks_and_stages_nothing(client, ed, db):
    from models import ImportBatch
    body = model(["Customer"], [{"de_name": "Customer", "field_name": "id"}])
    r = post(client, ed, "data-fields/check", body)
    assert r.status_code == 200, r.text
    assert r.json() == {"target": "data-fields", "sheet_name": "Data fields", "row_count": 1,
                        "file_warnings": r.json()["file_warnings"],
                        "other_sheets": [{"sheet": "Data entities", "target": "data-entities"}]}
    assert db.scalar(select(ImportBatch.import_batch_id).limit(1)) is None
    no_fields = edit_fields(body, lambda ws, col: setc(ws, col, 1, "ref_field_name", None))
    r = post(client, ed, "data-fields/check", no_fields)
    assert r.status_code == 422 and "ref_field_name" in r.json()["detail"]
    r = post(client, ed, "data-entities/check", model(None, [{"de_name": "Customer", "field_name": "id"}]))
    assert r.status_code == 422 and DE_TAG in r.json()["detail"]
    r = client.post(BULK.format(p=ed["p"], what="data-fields/check"), content=body,
                    headers={**ed["rv"], "Content-Type": xlsx.XLSX_MEDIA_TYPE})
    assert r.status_code == 403


# ---- values ---------------------------------------------------------------------------------

def test_sizes_types_and_yes_no_are_checked(client, ed, three):
    pv = stage(client, ed, "data-fields", model(None, [
        {"de_number": "DE-0001", "field_name": "a", "length": 5},
        {"de_number": "DE-0001", "field_name": "b", "data_type": "decimal", "precision": 4, "scale": 6},
        {"de_number": "DE-0001", "field_name": "c", "mandatory": "maybe", "pk_position": 40},
        {"de_number": "DE-0001", "field_name": "d", "data_type": "string", "length": "10.0"},
    ]))
    assert [codes_of(r) for r in pv["rows"]] == [["SIZE_NEEDS_TYPE"], ["SCALE_OVER_PRECISION"],
                                                 ["INVALID", "INVALID"], []]


# ---- the no-commit forms leave the existing field routes as they were ----------------------

def test_the_existing_field_routes_still_commit(client, ed, db, three):
    a = field(client, ed, three[0], "a")
    r = client.patch(f"{API}/data-fields/{a['data_field_id']}", headers=ed["h"],
                     json={"row_version": a["row_version"], "description": "x"})
    assert r.status_code == 200
    assert fields_of(db, ed["p"])["DE-0001/a"].description == "x"


# ---- sara M-1: swapping referenced fields inside one relationship -------------------------

def test_m1_a_swap_of_referenced_fields_in_one_relationship_is_a_validation_error(client, ed, db, three):
    cust, order = three[0], three[1]
    k1, k2 = pk(client, ed, cust, "id1", 1), pk(client, ed, cust, "id2", 2)
    field(client, ed, order, "a", ref_data_entity_id=cust["data_entity_id"], ref_data_field_id=k1["data_field_id"])
    field(client, ed, order, "b", ref_data_entity_id=cust["data_entity_id"], ref_data_field_id=k2["data_field_id"],
          fk_group=1)

    def swap(ws, col):
        for row in range(3, ws.max_row + 1):
            name = ws.cell(row, col["field_name"]).value
            if name in ("a", "b"):
                setc(ws, col, row, "ref_field_name", "id2" if name == "a" else "id1")
    pv = stage(client, ed, "data-fields", edit_fields(get(client, ed, "data-fields/export"), swap))
    codes = {r["business_key"]: codes_of(r) for r in pv["rows"]}
    assert codes == {"DE-0001/id1": [], "DE-0001/id2": [], "DE-0002/a": ["FK_SWAP"],
                     "DE-0002/b": ["FK_SWAP"]}, pv["rows"]
    assert commit(client, ed, pv).status_code == 422
    f = fields_of(db, ed["p"])
    assert (f["DE-0002/a"].ref_data_field_id, f["DE-0002/b"].ref_data_field_id) == \
        (k1["data_field_id"], k2["data_field_id"])


# ---- sara H-1/H-2: the import and the screen take turns on an entity -----------------------

@contextmanager
def _timed_session(ms: int = 200):
    """A session that gives up on a row lock after `ms` (55P03) instead of waiting for it."""
    s = SessionLocal()
    s.execute(text(f"SET lock_timeout = '{ms}ms'"))
    s.commit()                                  # a SET inside a rolled-back transaction would vanish
    try:
        yield s
    finally:
        s.rollback()
        s.execute(text("RESET lock_timeout"))
        s.commit()
        s.close()


@contextmanager
def _holding(sql: str, **params):
    """A side transaction holding a row lock until the block ends."""
    side = SessionLocal()
    try:
        side.execute(text(sql), params)
        yield side
    finally:
        side.rollback()
        side.close()


@pytest.mark.parametrize("held", ["parent", "referenced"])
def test_h2_the_imports_commit_waits_for_a_lock_on_an_entity_it_names(client, ed, db, world, three, held):
    """The commit holds every entity the sheet names (parent and FK target) FOR NO KEY UPDATE.
    A conflicting lock held elsewhere makes it wait; under lock_timeout that is an operational
    failure (503) that applies nothing and leaves the batch VALIDATED for a retry."""
    cust, order = three[0], three[1]
    k = pk(client, ed, cust, "id", 1)
    fk_ = field(client, ed, order, "cust_id", ref_data_entity_id=cust["data_entity_id"],
                ref_data_field_id=k["data_field_id"])                               # relationship 1
    pv = stage(client, ed, "data-fields", edit_fields(
        get(client, ed, "data-fields/export"), lambda ws, col: setc(ws, col, 4, "fk_group", "b")))
    assert [r["verdict"] for r in pv["rows"]] == ["MATCH", "UPDATE"], pv["rows"]
    batch_id, rv = pv["batch"]["import_batch_id"], pv["batch"]["row_version"]
    locked = (order if held == "parent" else cust)["data_entity_id"]

    with _holding("SELECT 1 FROM data_entity WHERE data_entity_id = :i FOR SHARE", i=locked), \
            _timed_session() as s:
        with pytest.raises(HTTPException) as e:
            bulk.commit(s, world["users"]["editor"], batch_id, rv, None)
        assert e.value.status_code == 503
    db.expire_all()
    assert preview(client, ed, pv)["batch"]["status"] == "VALIDATED"
    f = db.get(DataField, fk_["data_field_id"])
    assert (f.fk_group_no, f.row_version) == (1, fk_["row_version"])


@pytest.mark.parametrize("route", ["add", "change"])
def test_h1_a_screen_field_write_waits_while_an_import_holds_the_entity(client, ed, db, world, three, route):
    """sara H-1: seq_no and fk_group_no are max+1 reads; the screen must not write between an
    import's read and its write. With the import's lock held, the screen's write waits for it."""
    order = three[1]
    a = field(client, ed, order, "a")
    with _holding("SELECT 1 FROM data_entity WHERE data_entity_id = :i FOR NO KEY UPDATE",
                  i=order["data_entity_id"]), _timed_session() as s:
        with pytest.raises(OperationalError) as e:
            if route == "add":
                de_service.add_field(s, world["users"]["editor"], s.get(DataEntity, order["data_entity_id"]), "b", {})
            else:
                de_service.change_field(s, world["users"]["editor"], s.get(DataField, a["data_field_id"]), a["row_version"],
                                        None, {"description": "x"})
        assert e.value.orig.sqlstate == "55P03"
    f = fields_of(db, ed["p"])
    assert list(f) == ["DE-0002/a"] and f["DE-0002/a"].description is None
