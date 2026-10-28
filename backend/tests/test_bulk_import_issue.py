"""Slice 1e-a: the ISSUE import target (spec docs/slice1e_imports_spec.md §7.0, §7.3, round-2 N1/N4/N8; acceptance §10
items 1, 8, 9, 10, 12 for issues). Grade 2: a wrong issue row reaches a decision someone takes from the issue list but
is an audited, correctable update; added link ids are kept for a manual reversal. Every check runs on a real export,
edited the way a consultant would edit it.

Deliberately uncovered: two live stakeholders matching one name + unit (the live-name index per unit, I-16, makes it
unreachable except through a retired unit reusing a code — the RAISER_AMBIGUOUS branch is read, not run); two batches
over the same rows at once (row lock and row_version, proven in 4a)."""
import io

import openpyxl
from sqlalchemy import func, select

from models import Issue, IssueBr, Stakeholder
from tests.test_baseline import scope
from tests.test_br_fr import set_components
from tests.test_bulk_import import codes_of, commit
from tests.test_bulk_import_fr import by_key, by_row, put
from tests.test_issue_api import issue, resolve, stakeholder, unit
from tests.test_scope_api import API, ed  # noqa: F401 — ed is a fixture

URL = API + "/projects/{p}/bulk/issues/{what}"
TAG = "sheet:ISSUE"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def export(client, ed) -> bytes:
    r = client.get(URL.format(p=ed["p"], what="export"), headers=ed["h"])
    assert r.status_code == 200, r.text
    return r.content


def change(body: bytes, fn) -> bytes:
    """`fn(ws, col, rows, add)` on the issue sheet; rows keyed by issue number."""
    wb = openpyxl.load_workbook(io.BytesIO(body))
    ws = next(s for s in wb.worksheets if s.cell(1, 1).value == TAG)
    col = {ws.cell(1, i).value: i for i in range(2, ws.max_column + 1)}
    rows = {ws.cell(r, col["issue_number"]).value: r for r in range(3, ws.max_row + 1)}

    def add(**cells):
        at = max(ws.max_row, 2) + 1
        for k, v in cells.items():
            ws.cell(at, col[k]).value = v
    fn(ws, col, rows, add)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def staged(client, ed, body, expect=201):
    r = client.post(URL.format(p=ed["p"], what="validate"), content=body,
                    headers={**ed["h"], "Content-Type": XLSX, "X-VB-File-Name": "issues.xlsx"})
    assert r.status_code == expect, r.text
    return r.json()


def live(db, number):
    db.expire_all()
    return db.scalars(select(Issue).where(Issue.issue_number == number)).unique().one()


def stakeholders(db) -> int:
    return db.scalar(select(func.count()).select_from(Stakeholder))


NEW = {"title": "Tax code missing", "issue_type_code": "DATA", "severity_code": "HIGH", "raised_by": "Anan",
       "raised_by_org_unit": "SALES", "raised_date": "2026-10-20"}


# ---- §10.1: an unchanged export round-trips ------------------------------------------------------------------------

def test_an_unchanged_export_round_trips_as_all_match(client, ed, db):
    """P2 with D3: a closed issue, a linked BR, and a raiser whose name is also a stakeholder in another unit."""
    s = scope(client, ed)
    sales, fin = unit(client, ed, "SALES"), unit(client, ed, "FIN")
    anan = stakeholder(client, ed, "Anan", org_unit_id=sales["org_unit_id"])
    stakeholder(client, ed, "Anan", org_unit_id=fin["org_unit_id"])
    a = issue(client, ed, anan, raised_date="2026-10-01", target_date="2026-11-01")
    assert client.post(f"{API}/issues/{a['issue_id']}/brs", headers=ed["h"],
                       json={"br_id": s["br1"]["br_id"]}).status_code == 201
    closed = issue(client, ed, anan, "Closed one")
    assert resolve(client, ed, closed).status_code == 200
    pv = staged(client, ed, export(client, ed))
    rows = by_row(client, ed, pv)
    assert {r["verdict"] for r in rows} == {"MATCH"} and not any(r["errors"] for r in rows), rows


# ---- §10.8: the raiser (D3, N8) ------------------------------------------------------------------------------------

def test_a_raiser_matching_no_stakeholder_is_an_error_and_none_is_created(client, ed, db):
    scope(client, ed)
    sales = unit(client, ed, "SALES")
    stakeholder(client, ed, "Anan", org_unit_id=sales["org_unit_id"])
    before = stakeholders(db)
    body = change(export(client, ed), lambda ws, col, rows, add: (
        add(**{**NEW, "raised_by": "Somchai"}),
        add(**{**NEW, "title": "Blank unit", "raised_by_org_unit": None})))     # N8: blank = only the unit-less
    rows = by_row(client, ed, staged(client, ed, body))
    assert [codes_of(r) for r in rows] == [["RAISER_NOT_FOUND"], ["RAISER_NOT_FOUND"]]
    assert "Anan (SALES)" in rows[1]["errors"][0]["message"]                  # the candidates are named
    assert stakeholders(db) == before


def test_a_changed_raiser_resolves_by_name_and_unit_and_is_pinned(client, ed, db):
    scope(client, ed)
    sales, fin = unit(client, ed, "SALES"), unit(client, ed, "FIN")
    anan = stakeholder(client, ed, "Anan", org_unit_id=sales["org_unit_id"])
    other = stakeholder(client, ed, "Anan", org_unit_id=fin["org_unit_id"])
    a = issue(client, ed, anan)
    pv = staged(client, ed, change(export(client, ed), lambda ws, col, rows, add: put(
        ws, rows[a["issue_number"]], col["raised_by_org_unit"], "fin")))
    row = by_key(client, ed, pv)[a["issue_number"]]
    assert row["verdict"] == "UPDATE" and not row["errors"], row
    assert commit(client, ed, pv).status_code == 200
    assert live(db, a["issue_number"]).raised_by_stakeholder_id == other["stakeholder_id"]


# ---- §10.9: status ------------------------------------------------------------------------------------------------

def test_status_moves_among_open_statuses_only_and_never_on_a_closed_issue(client, ed, db):
    scope(client, ed)
    anan = stakeholder(client, ed, "Anan")
    a, closed = issue(client, ed, anan, "Open one"), issue(client, ed, anan, "Closed one")
    assert resolve(client, ed, closed).status_code == 200

    def edits(ws, col, rows, add):
        put(ws, rows[a["issue_number"]], col["status_code"], "RESOLVED")
        put(ws, rows[closed["issue_number"]], col["status_code"], "IN_PROGRESS")
    rows = by_key(client, ed, staged(client, ed, change(export(client, ed), edits)))
    assert codes_of(rows[a["issue_number"]]) == ["STATUS_CLOSES"]
    assert codes_of(rows[closed["issue_number"]]) == ["ISSUE_CLOSED"]

    pv = staged(client, ed, change(export(client, ed), lambda ws, col, rows, add: put(
        ws, rows[a["issue_number"]], col["status_code"], "IN_PROGRESS")))
    assert commit(client, ed, pv).status_code == 200
    assert live(db, a["issue_number"]).status_code == "IN_PROGRESS"


# ---- §10.12, N1, N4: new rows -------------------------------------------------------------------------------------

def test_a_new_row_is_raised_with_its_open_status_and_links(client, ed, db):
    s = scope(client, ed)
    sales = unit(client, ed, "SALES")
    anan = stakeholder(client, ed, "Anan", org_unit_id=sales["org_unit_id"])
    b1, b2 = s["br1"], s["br2"]
    body = change(export(client, ed), lambda ws, col, rows, add: (
        add(**NEW, status_code="IN_PROGRESS", target_date="2026-11-30",
            br_numbers=f"{b1['br_number']}; {b2['br_number']}"),
        add(title="No codes", raised_by="Anan", raised_by_org_unit="SALES", raised_date="2026-10-20")))
    rows = by_row(client, ed, staged(client, ed, body))
    assert not rows[0]["errors"] and codes_of(rows[1]) == ["REQUIRED", "REQUIRED"]       # N1: type and severity
    pv = staged(client, ed, change(export(client, ed), lambda ws, col, rows, add: add(
        **NEW, status_code="IN_PROGRESS", target_date="2026-11-30",
        br_numbers=f"{b1['br_number']}; {b2['br_number']}")))
    assert commit(client, ed, pv, ack=1).status_code == 200
    i = live(db, "ISS-0001")
    assert (i.title, i.status_code, i.raised_by_stakeholder_id, i.raised_date.isoformat(), i.target_date.isoformat()) \
        == ("Tax code missing", "IN_PROGRESS", anan["stakeholder_id"], "2026-10-20", "2026-11-30")    # N4
    assert set(db.scalars(select(IssueBr.br_id).where(IssueBr.issue_id == i.issue_id, IssueBr.is_active))) == \
        {b1["br_id"], b2["br_id"]}


def test_dates_are_checked_and_a_likely_duplicate_warns(client, ed):
    scope(client, ed)
    sales = unit(client, ed, "SALES")
    anan = stakeholder(client, ed, "Anan", org_unit_id=sales["org_unit_id"])
    a = issue(client, ed, anan, "Tax code missing", raised_date="2026-10-20")
    body = change(export(client, ed), lambda ws, col, rows, add: (
        add(**{**NEW, "title": " tax CODE missing "}),
        add(**{**NEW, "title": "Late", "target_date": "2026-10-01"}),
        add(**{**NEW, "title": "Bad date", "raised_date": "20/10/2026"})))
    rows = by_row(client, ed, staged(client, ed, body))
    new = [r for r in rows if r["verdict"] == "INSERT"]
    assert [w["code"] for w in new[0]["warnings"]] == ["POSSIBLE_DUPLICATE"]
    assert a["issue_number"] in new[0]["warnings"][0]["message"]
    assert codes_of(new[1]) == ["DATE_ORDER"]
    assert codes_of(new[2]) == ["INVALID"]


# ---- §7.0 links on an existing issue -------------------------------------------------------------------------------

def test_a_row_adding_a_br_links_it_and_keeps_the_link_id(client, ed, db):
    s = scope(client, ed)
    anan = stakeholder(client, ed, "Anan")
    a = issue(client, ed, anan)
    pv = staged(client, ed, change(export(client, ed), lambda ws, col, rows, add: put(
        ws, rows[a["issue_number"]], col["br_numbers"], s["br2"]["br_number"])))
    row = by_key(client, ed, pv)[a["issue_number"]]
    assert [(c["column"], c["after"]) for c in row["changes"]] == [("links_added", s["br2"]["br_number"])]
    assert commit(client, ed, pv).status_code == 200
    assert set(db.scalars(select(IssueBr.br_id).where(IssueBr.issue_id == a["issue_id"], IssueBr.is_active))) == \
        {s["br2"]["br_id"]}


# ---- §10.10: the component ----------------------------------------------------------------------------------------

def test_with_issue_off_template_export_and_validate_are_409(client, ed, world):
    scope(client, ed)
    body = export(client, ed)
    set_components(client, world, ed["p"], ["FUNCTION_REQUIREMENT", "TO_BE_DFD"])
    assert client.get(URL.format(p=ed["p"], what="export"), headers=ed["h"]).status_code == 409
    assert client.get(f"{API}/projects/{ed['p']}/bulk/templates/issues", headers=ed["h"]).status_code == 409
    staged(client, ed, body, expect=409)
