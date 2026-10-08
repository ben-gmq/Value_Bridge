"""Slice 4a-1 — the one import pipeline, first target Data Entities (docs/slice4a_spec.md §10,
rules 4a-R3…R18). Grade 1 for the commit path and upload safety: a silent failure here
overwrites a colleague's values or writes across projects, so each rule is proven directly
against the database, not through the response alone.

Test workbooks are built in code with openpyxl; the hostile ones are made by editing zip
members directly. Not covered here (4a-2/4a-3): fields, the two-batch data-model workbook
(criteria 9, 10, 20) and the purge (criterion 16)."""
import hashlib
import io
import threading
import zipfile
from dataclasses import replace
from urllib.parse import quote

import openpyxl
import pytest
from fastapi import HTTPException
from sqlalchemy import select, text
from sqlalchemy.exc import OperationalError

from database import SessionLocal
from models import AuditEvent, CodeMaster, DataEntity, ImportBatch, ImportRow
from services import bulk, xlsx
from services import data_entity as de_service
from tests.test_scope_api import API, ed, entity  # noqa: F401 — ed is a fixture

CODES = ("de_number", "de_name", "description", "business_owner_note", "row_token")
URL = API + "/projects/{p}/bulk/data-entities/{what}"
TAG = "sheet:DATA_ENTITY"


# ---- building and editing workbooks --------------------------------------------------------

def book(rows=(), codes=CODES, tag=TAG, title="Data entities", extra=()) -> bytes:
    """A workbook as a consultant's Excel would save it: row 1 codes, row 2 headers, data from 3.
    `rows` are dicts by column code; `extra` adds (title, A1) sheets."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.cell(1, 1, tag)
    for i, c in enumerate(codes, start=2):
        ws.cell(1, i, c)
        ws.cell(2, i, c.replace("_", " ").title())
    for r, row in enumerate(rows, start=3):
        for i, c in enumerate(codes, start=2):
            if row.get(c) is not None:
                ws.cell(r, i, row[c])
    for name, a1 in extra:
        wb.create_sheet(name).cell(1, 1, a1)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def edit(body: bytes, fn) -> bytes:
    """Open an export, let fn(ws, col) change it (col: code → column index), save it back."""
    wb = openpyxl.load_workbook(io.BytesIO(body))
    ws = next(s for s in wb.worksheets if s.cell(1, 1).value == TAG)
    col = {ws.cell(1, i).value: i for i in range(2, ws.max_column + 1)}
    fn(ws, col)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def cell(ws, col, sheet_row, code, value):
    ws.cell(sheet_row, col[code]).value = value


def rezip(body: bytes, change: dict) -> bytes:
    """Rewrite zip members: {name: fn(bytes) -> bytes}; a name not present is added from fn(b'')."""
    src = zipfile.ZipFile(io.BytesIO(body))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info)
            z.writestr(info.filename, change.pop(info.filename, lambda b: b)(data))
        for name, fn in change.items():
            z.writestr(name, fn(b""))
    return out.getvalue()


def with_dtd(xml: bytes, decl: bytes = b'<!DOCTYPE x [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>') -> bytes:
    head, sep, rest = xml.partition(b"?>")
    return head + sep + decl + rest if sep else decl + xml


# ---- calling the API ------------------------------------------------------------------------

def upload(client, ed, body, name="model.xlsx", headers=None, p=None):
    return client.post(URL.format(p=p or ed["p"], what="validate"), content=body,
                       headers={**(headers or ed["h"]), "Content-Type": xlsx.XLSX_MEDIA_TYPE,
                                "X-VB-File-Name": quote(name)})


def staged(client, ed, body, **kw):
    r = upload(client, ed, body, **kw)
    assert r.status_code == 201, r.text
    return r.json()


def export(client, ed, headers=None, p=None) -> bytes:
    r = client.get(URL.format(p=p or ed["p"], what="export"), headers=headers or ed["h"])
    assert r.status_code == 200, r.text
    return r.content


def commit(client, ed, pv, ack=None, headers=None, row_version=None):
    b = pv["batch"] if "batch" in pv else pv
    return client.post(f"{API}/bulk/batches/{b['import_batch_id']}/commit", headers=headers or ed["h"],
                       json={"row_version": row_version or b["row_version"], "acknowledged_inserts": ack})


def batch(client, ed, pv):
    b = pv["batch"] if "batch" in pv else pv
    return client.get(f"{API}/bulk/batches/{b['import_batch_id']}/preview", headers=ed["h"]).json()


def live(db, p):
    db.expire_all()
    return {d.de_number: d for d in db.scalars(select(DataEntity).where(DataEntity.project_id == p))}


def de_counter(db, p) -> int:
    return db.execute(text("SELECT last_value FROM project_sequence WHERE project_id = :p "
                           "AND sequence_code = 'DE'"), {"p": p}).scalar_one()


def codes_of(row) -> list[str]:
    return [e["code"] for e in row["errors"]]


@pytest.fixture
def three(client, ed):
    """Three live entities in the editor's project: DE-0001 Customer, DE-0002 Order, DE-0003 Product."""
    return [entity(client, ed, n) for n in ("Customer", "Order", "Product")]


# ---- criterion 2 (32, 111): the preview, and acknowledged inserts ------------------------------

def test_preview_shows_insert_update_match_with_before_after_and_clears(client, ed, three):
    client.patch(f"{API}/data-entities/{three[1]['data_entity_id']}", headers=ed["h"],
                 json={"row_version": three[1]["row_version"], "description": "One order"})

    def change(ws, col):
        cell(ws, col, 4, "de_name", "Sales order")            # DE-0002: rename and clear
        cell(ws, col, 4, "description", None)
        ws.cell(6, col["de_name"], "Invoice")                 # a new row, no number, no token
    pv = staged(client, ed, edit(export(client, ed), change))

    rows = {r["sheet_row_no"]: r for r in pv["rows"]}
    assert [rows[n]["verdict"] for n in (3, 4, 5, 6)] == ["MATCH", "UPDATE", "MATCH", "INSERT"]
    assert rows[4]["changes"] == [
        {"column": "de_name", "before": "Order", "after": "Sales order", "cleared": False},
        {"column": "description", "before": "One order", "after": None, "cleared": True}]
    assert rows[6]["changes"] == [{"column": "de_name", "before": None, "after": "Invoice", "cleared": False}]
    b = pv["batch"]
    assert (b["status"], b["insert_count"], b["update_count"], b["match_count"], b["error_count"]) == \
        ("VALIDATED", 1, 1, 2, 0)
    assert b["sheet_name"] == "Data entities" and b["file_name"] == "model.xlsx"


def test_inserts_need_the_count_acknowledged_even_when_every_row_is_an_insert(client, ed, db):
    pv = staged(client, ed, book([{"de_name": "Customer"}, {"de_name": "Order"}]))
    for ack in (None, 0, 1, 3):
        r = commit(client, ed, pv, ack=ack)
        assert r.status_code == 422 and "CREATE 2" in r.json()["detail"]
    assert live(db, ed["p"]) == {}
    assert batch(client, ed, pv)["batch"]["status"] == "VALIDATED"     # a typo never rejects the batch
    r = commit(client, ed, pv, ack=2)
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "COMMITTED"
    assert sorted((d.de_number, d.de_name) for d in live(db, ed["p"]).values()) == \
        [("DE-0001", "Customer"), ("DE-0002", "Order")]


# ---- criterion 1 (31): committing twice applies once -----------------------------------------

def test_committing_twice_applies_the_batch_once_and_the_second_call_is_409(client, ed, db):
    pv = staged(client, ed, book([{"de_name": "Customer"}]))
    assert commit(client, ed, pv, ack=1).status_code == 200
    again = commit(client, ed, pv, ack=1, row_version=batch(client, ed, pv)["batch"]["row_version"])
    assert again.status_code == 409
    assert list(live(db, ed["p"])) == ["DE-0001"]


def test_two_concurrent_commits_produce_one_set_of_rows(client, ed, db):
    pv = staged(client, ed, book([{"de_name": f"Entity {i}"} for i in range(20)]))
    b, user = pv["batch"], pv["batch"]["uploaded_by_user_id"]
    results, start = [], threading.Barrier(2)

    def run():
        with SessionLocal() as s:
            start.wait()
            try:
                bulk.commit(s, user, b["import_batch_id"], b["row_version"], 20)
                results.append(200)
            except HTTPException as e:
                results.append(e.status_code)
    threads = [threading.Thread(target=run) for _ in range(2)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert sorted(results) == [200, 409]
    assert len(live(db, ed["p"])) == 20 and de_counter(db, ed["p"]) == 20


# ---- criterion 3 (39): validate, retire, commit → rejected -------------------------------------

def test_a_de_retired_after_validate_rejects_the_batch_at_commit(client, ed, db, three):
    body = edit(export(client, ed), lambda ws, col: cell(ws, col, 3, "description", "Who buys"))
    pv = staged(client, ed, body)
    de = client.get(f"{API}/data-entities/{three[0]['data_entity_id']}", headers=ed["h"]).json()
    assert client.delete(f"{API}/data-entities/{de['data_entity_id']}", headers=ed["h"],
                         params={"row_version": de["row_version"]}).status_code == 204

    r = commit(client, ed, pv)
    assert r.status_code == 422, r.text
    after = batch(client, ed, pv)
    assert after["batch"]["status"] == "REJECTED" and after["batch"]["error_count"] == 1
    assert codes_of(next(x for x in after["rows"] if x["sheet_row_no"] == 3)) == ["ROW_RETIRED"]
    d = live(db, ed["p"])["DE-0001"]
    assert d.is_active is False and d.description is None              # nothing applied
    assert commit(client, ed, after).status_code == 409                # a REJECTED batch is final


# ---- criterion 4 (40): a colleague's edit since export -----------------------------------------

def test_a_row_edited_in_the_app_since_export_is_an_error_naming_the_current_value(client, ed, db, three):
    exported = export(client, ed)
    o = three[1]
    client.patch(f"{API}/data-entities/{o['data_entity_id']}", headers=ed["h"],
                 json={"row_version": o["row_version"], "description": "Colleague's text"})
    pv = staged(client, ed, edit(exported, lambda ws, col: cell(ws, col, 4, "description", "Mine")))
    row = next(r for r in pv["rows"] if r["sheet_row_no"] == 4)
    assert row["errors"] == [{"code": "ROW_STALE", "column": "description",
                              "message": "DE-0002 changed since export. Current value: “Colleague's text”",
                              "params": {"key": "DE-0002", "current": "Colleague's text"}}]
    assert row["verdict"] == "UPDATE" and pv["batch"]["error_count"] == 1
    r = commit(client, ed, pv)
    assert r.status_code == 422 and "Fix 1 rows" in r.json()["detail"]
    assert live(db, ed["p"])["DE-0002"].description == "Colleague's text"

    fixed = staged(client, ed, edit(export(client, ed), lambda ws, col: cell(ws, col, 4, "description", "Mine")))
    assert fixed["batch"]["error_count"] == 0
    assert commit(client, ed, fixed).status_code == 200
    assert live(db, ed["p"])["DE-0002"].description == "Mine"


# ---- criterion 5: a token whose row is gone is never an INSERT; stale at commit ------------------

def test_a_token_whose_entity_was_retired_is_an_error_never_an_insert(client, ed, db, three):
    exported = export(client, ed)
    p = three[2]
    client.delete(f"{API}/data-entities/{p['data_entity_id']}", headers=ed["h"],
                  params={"row_version": p["row_version"]})
    pv = staged(client, ed, exported)
    row = next(r for r in pv["rows"] if r["sheet_row_no"] == 5)
    assert row["verdict"] == "UPDATE" and codes_of(row) == ["ROW_RETIRED"]
    assert pv["batch"]["insert_count"] == 0


def test_a_row_going_stale_between_validate_and_commit_applies_nothing(client, ed, db, three):
    body = edit(export(client, ed), lambda ws, col: (cell(ws, col, 3, "description", "Buys"),
                                                     cell(ws, col, 5, "description", "Sold")))
    pv = staged(client, ed, body)
    p = three[2]
    client.patch(f"{API}/data-entities/{p['data_entity_id']}", headers=ed["h"],
                 json={"row_version": p["row_version"], "de_name": "Product line"})

    r = commit(client, ed, pv)
    assert r.status_code == 422
    after = batch(client, ed, pv)
    assert after["batch"]["status"] == "REJECTED"
    stale = next(x for x in after["rows"] if x["sheet_row_no"] == 5)
    assert [(e["code"], e["column"], e["params"]["current"]) for e in stale["errors"]] == [
        ("ROW_STALE", "de_name", "Product line"), ("ROW_STALE", "description", None)]
    d = live(db, ed["p"])
    assert d["DE-0001"].description is None and d["DE-0003"].description is None    # row 3 not applied either


# ---- criterion 6: another project's export; a partial sort -------------------------------------

def test_an_export_from_another_project_is_refused_whole(client, ed, world, db):
    client.post(f"{API}/projects/{world['p_beta']}/access", headers=world["admin"],
                json={"user_id": world["users"]["owner"], "project_role_code": "EDITOR"})
    other = {"h": world["admin"], "p": world["p_beta"]}
    entity(client, other, "Ledger")
    r = upload(client, ed, export(client, other))
    assert r.status_code == 422 and r.json()["detail"] == "This file was exported from another project."
    assert client.get(f"{API}/projects/{ed['p']}/data-entities", headers=ed["h"]).json() == []
    assert db.scalar(select(ImportBatch.import_batch_id)) is None           # nothing staged


def test_a_partial_sort_that_shifts_numbers_against_tokens_is_row_errors_not_overwrites(client, ed, db, three):
    def sort_numbers_only(ws, col):                 # rows 3 and 5 swap number and name, not the token
        a = [ws.cell(3, col[c]).value for c in ("de_number", "de_name")]
        b = [ws.cell(5, col[c]).value for c in ("de_number", "de_name")]
        for c, v in zip(("de_number", "de_name"), b):
            cell(ws, col, 3, c, v)
        for c, v in zip(("de_number", "de_name"), a):
            cell(ws, col, 5, c, v)
    pv = staged(client, ed, edit(export(client, ed), sort_numbers_only))
    rows = {r["sheet_row_no"]: r for r in pv["rows"]}
    assert codes_of(rows[3]) == ["TOKEN_ROW_MISMATCH"] and codes_of(rows[5]) == ["TOKEN_ROW_MISMATCH"]
    assert rows[3]["errors"][0]["params"] == {"key": "DE-0003", "token_key": "DE-0001"}
    assert rows[4]["verdict"] == "MATCH"
    assert commit(client, ed, pv).status_code == 422
    assert {n: d.de_name for n, d in live(db, ed["p"]).items()} == \
        {"DE-0001": "Customer", "DE-0002": "Order", "DE-0003": "Product"}


# ---- criterion 7: blank clears; a missing column; a formula with no cached value ----------------

def test_a_blank_description_clears_and_a_missing_column_is_422(client, ed, db, three):
    c = three[0]
    client.patch(f"{API}/data-entities/{c['data_entity_id']}", headers=ed["h"],
                 json={"row_version": c["row_version"], "description": "Who buys"})
    pv = staged(client, ed, edit(export(client, ed), lambda ws, col: cell(ws, col, 3, "description", None)))
    assert commit(client, ed, pv).status_code == 200
    assert live(db, ed["p"])["DE-0001"].description is None

    r = upload(client, ed, book([{"de_name": "X"}], codes=tuple(c for c in CODES if c != "description")))
    assert r.status_code == 422 and r.json()["detail"].startswith("The sheet is missing these columns: description")


def test_a_formula_with_no_cached_value_is_a_row_error(client, ed):
    pv = staged(client, ed, book([{"de_name": "Customer", "description": "=A1&\"x\""}]))
    [row] = pv["rows"]
    assert codes_of(row) == ["FORMULA_NO_VALUE"] and row["errors"][0]["column"] == "description"


# ---- criterion 8: a round trip with no edits; the escape set --------------------------------------

def test_export_then_reimport_unchanged_is_all_match_and_writes_nothing(client, ed, db, three):
    for d, text_ in zip(three, ("-1 minus first", "=SUM(A1)", "@user +plus")):
        client.patch(f"{API}/data-entities/{d['data_entity_id']}", headers=ed["h"],
                     json={"row_version": d["row_version"], "description": text_})
    before = {n: (d.row_version, d.updated_at) for n, d in live(db, ed["p"]).items()}
    exported = export(client, ed)

    ws = openpyxl.load_workbook(io.BytesIO(exported))["Data entities"]
    desc = ws.cell(3, 4)
    assert (desc.value, desc.data_type, desc.quotePrefix) == ("-1 minus first", "s", True)
    assert (ws.cell(4, 4).value, ws.cell(4, 4).data_type) == ("=SUM(A1)", "s")   # text, not a formula
    assert ws.row_dimensions[1].hidden and ws.cell(1, 1).value == TAG and ws.protection.sheet

    pv = staged(client, ed, exported)
    assert [r["verdict"] for r in pv["rows"]] == ["MATCH"] * 3
    assert (pv["batch"]["insert_count"], pv["batch"]["update_count"]) == (0, 0)
    assert commit(client, ed, pv).status_code == 200
    after = live(db, ed["p"])
    assert {n: (d.row_version, d.updated_at) for n, d in after.items()} == before
    assert after["DE-0001"].description == "-1 minus first"


def _with_types(db, p, codes):
    for i, code in enumerate(codes):
        db.add(CodeMaster(project_id=p, category="FIELD_DATA_TYPE", code=code, label=code, sort_order=100 + i))
    db.commit()


def test_the_template_lists_this_projects_codes_and_a_long_list_sits_on_the_hidden_lookup_sheet(db, ed, world):
    """Criteria 15 (47) and 8 (4a-R9): the lists come from this project's resolved codes, and
    every lookup cell goes through write_cell — a value starting with = stays text."""
    typed = replace(bulk.DATA_ENTITY, columns=bulk.DE_COLUMNS + (
        bulk.Column("data_type", "Data type", category="FIELD_DATA_TYPE"),))
    _with_types(db, ed["p"], ["=HYPERLINK", "GEOMETRY"])
    _with_types(db, world["p_fin"], ["ONLY_IN_FIN"])

    wb = openpyxl.load_workbook(io.BytesIO(bulk.template(db, ed["p"], typed)))
    [dv] = wb["Data entities"].data_validations.dataValidation
    listed = dv.formula1.strip('"').split(",")
    assert "GEOMETRY" in listed and "=HYPERLINK" in listed and "STRING" in listed
    assert "ONLY_IN_FIN" not in listed

    _with_types(db, ed["p"], [f"LONG_TYPE_CODE_NUMBER_{i:02d}" for i in range(12)])
    wb = openpyxl.load_workbook(io.BytesIO(bulk.template(db, ed["p"], typed)))
    [dv] = wb["Data entities"].data_validations.dataValidation
    lookup = wb[xlsx.LOOKUP_SHEET]
    assert lookup.sheet_state == "hidden" and dv.formula1.startswith(f"'{xlsx.LOOKUP_SHEET}'!$A$2:$A$")
    values = {lookup.cell(r, 1).value: lookup.cell(r, 1) for r in range(2, lookup.max_row + 1)}
    evil = values["=HYPERLINK"]
    assert (evil.data_type, evil.quotePrefix) == ("s", True)
    assert "LONG_TYPE_CODE_NUMBER_11" in values


# ---- criterion 11: a failure inside a service; a helper that commits ------------------------------

def test_a_service_failure_on_the_last_row_changes_nothing_and_burns_no_number(client, ed, db, monkeypatch):
    pv = staged(client, ed, book([{"de_name": f"E{i}"} for i in range(3)]))
    real, calls = de_service.add_entity, []

    def failing(db_, *a, **k):
        calls.append(1)
        if len(calls) == 3:
            raise HTTPException(422, "refused by the service")
        return real(db_, *a, **k)
    monkeypatch.setattr(de_service, "add_entity", failing)

    r = commit(client, ed, pv, ack=3)
    assert r.status_code == 422 and "Row 5 could not be applied: refused by the service" in r.json()["detail"]
    assert live(db, ed["p"]) == {} and de_counter(db, ed["p"]) == 0
    after = batch(client, ed, pv)
    assert after["batch"]["status"] == "REJECTED"
    assert codes_of(next(x for x in after["rows"] if x["sheet_row_no"] == 5)) == ["APPLY_FAILED"]


def test_a_helper_that_commits_during_the_apply_raises_and_nothing_is_applied(client, ed, db, monkeypatch):
    pv = staged(client, ed, book([{"de_name": "Customer"}, {"de_name": "Order"}]))
    real = de_service.add_entity

    def committing(db_, *a, **k):
        de = real(db_, *a, **k)
        db_.commit()
        return de
    monkeypatch.setattr(de_service, "add_entity", committing)
    with pytest.raises(RuntimeError, match="Session.commit called during an import's apply"):
        commit(client, ed, pv, ack=2)
    assert live(db, ed["p"]) == {} and de_counter(db, ed["p"]) == 0
    assert batch(client, ed, pv)["batch"]["status"] == "REJECTED"


# ---- criterion 12: operational vs rule failures (4a-R7) ------------------------------------------

class _Deadlock(Exception):
    sqlstate = "40P01"


def test_an_operational_failure_leaves_the_batch_validated_and_a_retry_succeeds(client, ed, db, monkeypatch):
    pv = staged(client, ed, book([{"de_name": "Customer"}]))
    real = de_service.add_entity

    def deadlocked(*a, **k):
        raise OperationalError("INSERT …", {}, _Deadlock())
    monkeypatch.setattr(de_service, "add_entity", deadlocked)
    r = commit(client, ed, pv, ack=1)
    assert r.status_code == 503
    assert batch(client, ed, pv)["batch"]["status"] == "VALIDATED" and live(db, ed["p"]) == {}

    monkeypatch.setattr(de_service, "add_entity", real)
    assert commit(client, ed, pv, ack=1).status_code == 200
    assert list(live(db, ed["p"])) == ["DE-0001"]
    kinds = [e.detail["kind"] for e in db.scalars(select(AuditEvent).where(
        AuditEvent.event_type == "IMPORT_COMMIT_FAILED"))]
    assert kinds == ["OPERATIONAL"]                                          # every failure audited


def test_a_rule_failure_never_overwrites_a_racing_successful_commit(client, ed, db):
    pv = staged(client, ed, book([{"de_name": "Customer"}]))
    assert commit(client, ed, pv, ack=1).status_code == 200
    status = bulk._statuses(db)
    assert bulk._reject(db, pv["batch"]["import_batch_id"], status, {3: [bulk.msg("X", "x")]}) is False
    db.commit()
    after = batch(client, ed, pv)
    assert after["batch"]["status"] == "COMMITTED" and after["rows"][0]["is_valid"] is True


def test_every_failed_commit_attempt_is_audited(client, ed, db):
    pv = staged(client, ed, book([{"de_name": "Customer"}]))
    commit(client, ed, pv, ack=5)
    [ev] = db.scalars(select(AuditEvent).where(AuditEvent.event_type == "IMPORT_COMMIT_FAILED"))
    assert ev.detail["kind"] == "REFUSED" and ev.target_id == pv["batch"]["import_batch_id"]


# ---- criterion 13 (75): hostile files ------------------------------------------------------------

PARTS = ("xl/workbook.xml", "xl/sharedStrings.xml", "xl/worksheets/sheet1.xml")


def _shared_strings_book() -> bytes:
    """openpyxl writes inline strings; give the file a real shared-string table to attack."""
    sst = (b'<?xml version="1.0" encoding="UTF-8"?><sst xmlns="http://schemas.openxmlformats.org/'
           b'spreadsheetml/2006/main" count="0" uniqueCount="0"></sst>')
    override = (b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-'
                b'officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    body = book([{"de_name": "Customer"}])
    return rezip(body, {"xl/sharedStrings.xml": lambda b: b or sst,
                        "[Content_Types].xml": lambda b: b.replace(b"</Types>", override)})


@pytest.mark.parametrize("part", PARTS)
def test_a_dtd_or_external_entity_in_any_part_is_422_never_resolved(client, ed, part, monkeypatch):
    hostile = rezip(_shared_strings_book(), {part: with_dtd})
    r = upload(client, ed, hostile)
    assert r.status_code == 422 and "document type declaration" in r.json()["detail"]

    # The second gate on its own: with the pre-pass switched off, the parser itself refuses.
    monkeypatch.setattr(xlsx, "zip_prepass", lambda body: None)
    with pytest.raises(HTTPException) as e:
        xlsx.read_rows(hostile, "DATA_ENTITY", CODES)
    assert e.value.status_code == 422 and "entity declaration" in e.value.detail


def test_the_parser_is_hardened_by_defusedxml():
    from defusedxml import EntitiesForbidden
    from openpyxl.xml.functions import fromstring
    with pytest.raises(EntitiesForbidden):
        fromstring(b'<!DOCTYPE x [<!ENTITY e "boom">]><x>&e;</x>')


def test_a_non_utf8_xml_part_is_refused(client, ed):
    body = rezip(book([{"de_name": "C"}]), {"xl/workbook.xml": lambda b: b'<?xml version="1.0" '
                                                                         b'encoding="cp500"?>' + b})
    assert upload(client, ed, body).status_code == 422


def test_a_zip_bomb_and_an_oversized_text_table_are_refused_in_the_pre_pass(client, ed):
    bomb = rezip(book([{"de_name": "C"}]), {"xl/media/pad.bin": lambda b: b"\0" * (101 * xlsx.MB)})
    assert len(bomb) < 1 * xlsx.MB
    r = upload(client, ed, bomb)
    assert r.status_code == 422 and "100 MB" in r.json()["detail"]

    big_sst = rezip(_shared_strings_book(), {"xl/sharedStrings.xml": lambda b: b + b" " * (21 * xlsx.MB)})
    r = upload(client, ed, big_sst)
    assert r.status_code == 422 and "20 MB" in r.json()["detail"]


def test_a_text_table_or_xml_part_under_another_name_is_still_checked(client, ed):
    """sara MEDIUM-2: openpyxl finds parts by the package's own names, so the pre-pass does too."""
    renamed = rezip(_shared_strings_book(), {
        "[Content_Types].xml": lambda b: b.replace(b"/xl/sharedStrings.xml", b"/xl/s.bin")})
    renamed = rezip(renamed, {"xl/s.bin": lambda b: b"", "xl/sharedStrings.xml": lambda b: b})
    big = rezip(renamed, {"xl/s.bin": lambda b: b'<sst xmlns="x">' + b" " * (21 * xlsx.MB) + b"</sst>"})
    r = upload(client, ed, big)
    assert r.status_code == 422 and "20 MB" in r.json()["detail"]
    hostile = rezip(book([{"de_name": "C"}]), {"xl/media/odd.bin": lambda b: with_dtd(b"<x/>")})
    r = upload(client, ed, hostile)
    assert r.status_code == 422 and "document type declaration" in r.json()["detail"]


def test_a_repeated_or_backwards_row_number_is_422_not_500(client, ed):
    """sara MEDIUM-1: row numbers come from the file, so they are checked, never trusted."""
    body = book([{"de_name": "Customer"}, {"de_name": "Order"}])
    hostile = rezip(body, {"xl/worksheets/sheet1.xml": lambda b: b.replace(b'<row r="4"', b'<row r="3"')})
    r = upload(client, ed, hostile)
    assert r.status_code == 422 and "out of order" in r.json()["detail"]


def test_more_than_5000_rows_or_200_columns_is_refused_mid_stream(client, ed):
    r = upload(client, ed, book([{"de_name": f"E{i}"} for i in range(xlsx.MAX_ROWS + 1)]))
    assert r.status_code == 422 and "more than 5,000 rows" in r.json()["detail"]

    def wide(ws, col):
        ws.cell(3, xlsx.MAX_COLS + 1, "x")
    r = upload(client, ed, edit(book([{"de_name": "C"}]), wide))
    assert r.status_code == 422 and "more than 200 columns" in r.json()["detail"]


def test_a_third_concurrent_upload_is_429(client, ed):
    with xlsx.parse_slot(), xlsx.parse_slot():
        r = upload(client, ed, book([{"de_name": "C"}]))
    assert r.status_code == 429
    assert upload(client, ed, book([{"de_name": "C"}])).status_code == 201     # the slots come back


def test_a_file_that_is_not_a_workbook_is_422(client, ed):
    assert upload(client, ed, b"PK\x03\x04 not really").status_code == 422
    assert upload(client, ed, b"plain text").status_code == 422


# ---- criterion 14 (138, 152): body limits and no spooling --------------------------------------

def test_a_50mb_body_is_413_before_auth_or_the_handler(client, ed, monkeypatch):
    called = []
    monkeypatch.setattr(bulk, "validate", lambda *a, **k: called.append(1))
    r = client.post(URL.format(p=ed["p"], what="validate"), content=b"\0" * (50 * xlsx.MB),
                    headers={"Content-Type": xlsx.XLSX_MEDIA_TYPE})            # no JWT: 413, not 401
    assert r.status_code == 413 and called == []


def test_a_chunked_over_cap_body_is_413(client, ed, monkeypatch):
    called = []
    monkeypatch.setattr(bulk, "validate", lambda *a, **k: called.append(1))

    def chunks():
        for _ in range(11):
            yield b"\0" * xlsx.MB
    r = client.post(URL.format(p=ed["p"], what="validate"), content=chunks(),
                    headers={**ed["h"], "Content-Type": xlsx.XLSX_MEDIA_TYPE})
    assert "content-length" not in {k.lower() for k in r.request.headers}
    assert r.status_code == 413 and called == []


def _exactly(size: int) -> bytes:
    """A real workbook padded to exactly `size` bytes with a stored (uncompressed) member."""
    src = zipfile.ZipFile(io.BytesIO(book([{"de_name": "Customer"}])))

    def build(pad: int) -> bytes:
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            for info in src.infolist():
                z.writestr(info.filename, src.read(info))
            z.writestr("xl/media/pad.bin", b"\0" * pad, compress_type=zipfile.ZIP_STORED)
        return buf.getvalue()
    out = build(size - len(build(0)))
    assert len(out) == size
    return out


def test_exactly_10mb_is_accepted_and_nothing_is_spooled_to_disk(client, ed, monkeypatch):
    import tempfile

    def no_disk(*a, **k):
        raise AssertionError("an upload touched the disk")
    for name in ("TemporaryFile", "NamedTemporaryFile", "SpooledTemporaryFile"):
        monkeypatch.setattr(tempfile, name, no_disk)
    r = upload(client, ed, _exactly(10 * xlsx.MB))
    assert r.status_code == 201, r.text
    r = upload(client, ed, _exactly(10 * xlsx.MB + 1))
    assert r.status_code == 413


def test_the_upload_must_be_the_raw_xlsx_body(client, ed):
    r = client.post(URL.format(p=ed["p"], what="validate"), files={"file": ("a.xlsx", book())},
                    headers=ed["h"])
    assert r.status_code == 415


# ---- criterion 15: a rejected batch burns no number ----------------------------------------------

def test_a_batch_with_errors_burns_no_de_number(client, ed, db):
    bad = staged(client, ed, book([{"de_name": "Customer"}, {"de_name": "x" * 201}]))
    assert bad["batch"]["error_count"] == 1
    assert commit(client, ed, bad, ack=2).status_code == 422
    assert de_counter(db, ed["p"]) == 0
    good = staged(client, ed, book([{"de_name": "Customer"}]))
    assert commit(client, ed, good, ack=1).status_code == 200
    assert list(live(db, ed["p"])) == ["DE-0001"]


# ---- criterion 17 (Q5): before-values; 4a-R18: the file's SHA-256 ------------------------------------

def test_commit_stores_before_values_and_audits_the_file_hash(client, ed, db, three):
    c = three[0]
    client.patch(f"{API}/data-entities/{c['data_entity_id']}", headers=ed["h"],
                 json={"row_version": c["row_version"], "description": "Who buys"})
    body = edit(export(client, ed), lambda ws, col: (cell(ws, col, 3, "description", "Who pays"),
                                                    ws.cell(6, col["de_name"], "Invoice")))
    pv = staged(client, ed, body)
    assert commit(client, ed, pv, ack=1).status_code == 200

    rows = {r.sheet_row_no: r for r in db.scalars(select(ImportRow))}
    assert rows[3].payload["before"] == {"de_name": "Customer", "description": "Who buys",
                                         "business_owner_note": None, "row_version": 2}
    assert "before" not in rows[4].payload and "before" not in rows[6].payload     # MATCH, INSERT
    assert rows[3].committed_target_id == c["data_entity_id"]
    assert rows[4].committed_target_id == three[1]["data_entity_id"]
    assert rows[6].committed_target_id == live(db, ed["p"])["DE-0004"].data_entity_id

    [ev] = db.scalars(select(AuditEvent).where(AuditEvent.event_type == "IMPORT_COMMITTED"))
    assert ev.detail["file_sha256"] == hashlib.sha256(body).hexdigest()
    assert (ev.detail["entity"], ev.detail["insert_count"], ev.detail["update_count"]) == ("DATA_ENTITY", 1, 1)
    assert ev.target_id == pv["batch"]["import_batch_id"]

    shown = {r["sheet_row_no"]: r for r in batch(client, ed, pv)["rows"]}
    assert shown[3]["changes"] == [{"column": "description", "before": "Who buys", "after": "Who pays",
                                    "cleared": False}]


# ---- criterion 18: roles and visibility --------------------------------------------------------

def test_a_reviewer_can_export_but_not_template_validate_preview_or_commit(client, ed, three):
    assert export(client, ed, headers=ed["rv"])
    assert client.get(f"{API}/projects/{ed['p']}/bulk/templates/data-entities",
                      headers=ed["rv"]).status_code == 403
    assert upload(client, ed, book([{"de_name": "X"}]), headers=ed["rv"]).status_code == 403
    pv = staged(client, ed, book([{"de_name": "X"}]))
    b = pv["batch"]["import_batch_id"]
    assert client.get(f"{API}/bulk/batches/{b}/preview", headers=ed["rv"]).status_code == 403
    assert commit(client, ed, pv, ack=1, headers=ed["rv"]).status_code == 403


def test_a_batch_in_a_project_you_cannot_see_is_404(client, ed, world):
    pv = staged(client, ed, book([{"de_name": "X"}]))
    b = pv["batch"]["import_batch_id"]
    for h in (ed["out"],):
        assert client.get(f"{API}/bulk/batches/{b}", headers=h).status_code == 404
        assert commit(client, ed, pv, ack=1, headers=h).status_code == 404
    assert upload(client, ed, book([{"de_name": "X"}]), headers=ed["out"]).status_code == 404
    assert client.get(URL.format(p=ed["p"], what="export"), headers=ed["out"]).status_code == 404
    assert client.get(f"{API}/bulk/batches/999999", headers=ed["h"]).status_code == 404


def test_the_template_and_unknown_targets(client, ed):
    r = client.get(f"{API}/projects/{ed['p']}/bulk/templates/data-entities", headers=ed["h"])
    assert r.status_code == 200 and r.headers["content-type"] == xlsx.XLSX_MEDIA_TYPE
    assert r.headers["content-disposition"].startswith('attachment; filename="vb-data-entities-template-p')
    read = xlsx.read_rows(r.content, "DATA_ENTITY", CODES)
    assert read.rows == [] and read.sheet_name == "Data entities"
    for target in ("data-fields", "issues"):
        assert client.get(f"{API}/projects/{ed['p']}/bulk/templates/{target}", headers=ed["h"]).status_code == 404
        assert client.get(f"{API}/projects/{ed['p']}/bulk/{target}/export", headers=ed["h"]).status_code == 404
        r = client.post(f"{API}/projects/{ed['p']}/bulk/{target}/validate", content=book([{"de_name": "X"}]),
                        headers={**ed["h"], "Content-Type": xlsx.XLSX_MEDIA_TYPE})
        assert r.status_code == 404


# ---- 4a-R11: exactly one sheet per target ------------------------------------------------------

def test_no_tagged_sheet_or_two_is_422_naming_the_sheets_cut_to_31(client, ed):
    long = "A sheet whose name runs on and on"[:31]
    r = upload(client, ed, book([{"de_name": "X"}], tag="something else", title=long))
    assert r.status_code == 422
    assert f"'{long}'" in r.json()["detail"] and "No sheet is marked sheet:DATA_ENTITY" in r.json()["detail"]

    r = upload(client, ed, book([{"de_name": "X"}], extra=[("Copy", TAG)]))
    assert r.status_code == 422 and "'Data entities', 'Copy'" in r.json()["detail"]
    assert xlsx.sheet_label("x" * 40) == "'" + "x" * 31 + "'"


def test_an_empty_sheet_and_unknown_columns(client, ed):
    assert upload(client, ed, book([])).status_code == 422
    pv = staged(client, ed, book([{"de_name": "X", "notes": "ignored"}], codes=CODES + ("notes",)))
    assert [w["code"] for w in pv["batch"]["file_warnings"]] == ["COLUMN_IGNORED"]


# ---- the entity rules: R4, names, keys, lengths -------------------------------------------------

def test_a_number_without_a_token_is_export_first(client, ed, three):
    pv = staged(client, ed, book([{"de_number": "de-0001", "de_name": "Customer"}]))
    [row] = pv["rows"]
    assert row["verdict"] == "UPDATE" and codes_of(row) == ["EXPORT_FIRST"]
    assert row["business_key"] == "DE-0001"


def test_a_tampered_or_foreign_key_token_is_invalid(client, ed, three, monkeypatch):
    def tamper(ws, col):
        tok = ws.cell(3, col["row_token"]).value
        cell(ws, col, 3, "row_token", tok.replace(".data_entity.", ".data_entity.9", 1))
    pv = staged(client, ed, edit(export(client, ed), tamper))
    assert codes_of(next(r for r in pv["rows"] if r["sheet_row_no"] == 3)) == ["TOKEN_INVALID"]

    from config import get_settings
    exported = export(client, ed)
    monkeypatch.setattr(get_settings(), "row_token_key", "another-key-" + "x" * 40)
    pv = staged(client, ed, exported)
    assert all(codes_of(r) == ["TOKEN_INVALID"] for r in pv["rows"])


def test_a_genuine_token_for_another_table_or_a_blank_number_is_an_error(client, ed, three):
    d = three[0]
    field_token = bulk.row_token(ed["p"], "data_field", d["data_entity_id"], d["row_version"])
    pv = staged(client, ed, edit(export(client, ed), lambda ws, col: (
        cell(ws, col, 3, "row_token", field_token), cell(ws, col, 4, "de_number", None))))
    rows = {r["sheet_row_no"]: r for r in pv["rows"]}
    assert codes_of(rows[3]) == ["TOKEN_OTHER_TABLE"]
    assert rows[4]["verdict"] == "INSERT" and codes_of(rows[4]) == ["TOKEN_ROW_MISMATCH", "NAME_EXISTS"]


def test_the_token_reads_back_its_claims_and_rejects_any_edit():
    tok = bulk.row_token(7, "data_entity", 412, 3)
    assert tok.startswith("v1.7.data_entity.412.3.")
    assert bulk.read_token(tok) == bulk.Claims(7, "data_entity", 412, 3)
    for bad in (tok.replace(".3.", ".4."), tok[:-1] + ("A" if tok[-1] != "A" else "B"), "", "v1.7"):
        assert bulk.read_token(bad) is None


def test_names_new_renamed_swapped_and_duplicated(client, ed, three):
    def change(ws, col):
        cell(ws, col, 3, "de_name", "Order")              # rename onto a live name → error
        cell(ws, col, 4, "de_name", "Customer")           # the other half of the swap → error
        ws.cell(6, col["de_name"], "product ")            # a new row onto a live name (case, space)
        ws.cell(7, col["de_name"], "Invoice")
        ws.cell(8, col["de_name"], "INVOICE")             # a duplicate inside the file
    pv = staged(client, ed, edit(export(client, ed), change))
    rows = {r["sheet_row_no"]: r for r in pv["rows"]}
    assert codes_of(rows[3]) == ["NAME_TAKEN"] and rows[3]["errors"][0]["params"] == {"key": "DE-0002"}
    assert codes_of(rows[4]) == ["NAME_TAKEN"]
    assert codes_of(rows[6]) == ["NAME_EXISTS"] and "DE-0003" in rows[6]["errors"][0]["message"]
    assert codes_of(rows[7]) == [] and codes_of(rows[8]) == ["DUPLICATE_NAME"]


def test_a_case_only_rename_is_an_update(client, ed, db, three):
    pv = staged(client, ed, edit(export(client, ed), lambda ws, col: cell(ws, col, 3, "de_name", "CUSTOMER")))
    assert [r["verdict"] for r in pv["rows"]][0] == "UPDATE" and pv["batch"]["error_count"] == 0
    assert commit(client, ed, pv).status_code == 200
    assert live(db, ed["p"])["DE-0001"].de_name == "CUSTOMER"


def test_duplicate_numbers_lengths_and_a_blank_name(client, ed, three):
    def change(ws, col):
        for c in CODES:
            ws.cell(6, col[c], ws.cell(3, col[c]).value)          # row 3 repeated
        cell(ws, col, 4, "description", "d" * 4001)
        cell(ws, col, 5, "de_name", None)
    pv = staged(client, ed, edit(export(client, ed), change))
    rows = {r["sheet_row_no"]: r for r in pv["rows"]}
    assert codes_of(rows[6]) == ["DUPLICATE_KEY"] and rows[6]["business_key"] is None
    assert codes_of(rows[4]) == ["TOO_LONG"] and rows[4]["errors"][0]["params"] == {"max": 4000, "length": 4001}
    assert codes_of(rows[5]) == ["REQUIRED"]


# ---- 4a-R14: names are bound to the batch's project ----------------------------------------------

def test_a_name_live_only_in_another_project_never_leaks(client, ed, world):
    other = {"h": world["admin"], "p": world["p_beta"]}
    for n in ("Ledger", "Journal", "Customer"):
        entity(client, other, n)                                   # Customer is DE-0003 over there
    pv = staged(client, ed, book([{"de_name": "Customer"}]))
    [row] = pv["rows"]
    assert row["verdict"] == "INSERT" and row["errors"] == []
    r = upload(client, ed, book([{"de_name": "Customer"}]))
    assert "DE-0003" not in r.text and "DE-000" not in r.text


# ---- 4a-R15: whitespace -------------------------------------------------------------------------

def test_unicode_whitespace_is_stripped_before_validation_and_lookup(client, ed, db, three):
    pv = staged(client, ed, book([{"de_name": " Customer "}, {"de_name": " ​Ledger  "}]))
    rows = pv["rows"]
    assert codes_of(rows[0]) == ["NAME_EXISTS"]
    assert rows[1]["changes"][0]["after"] == "Ledger"

    def padded(ws, col):
        cell(ws, col, 3, "de_name", "Customer  ")
    pv = staged(client, ed, edit(export(client, ed), padded))
    assert pv["rows"][0]["verdict"] == "MATCH"
    assert xlsx.clean_cell(" \t x 　") == "x" and xlsx.clean_cell("   ") is None


# ---- 4a-R8: the display name -------------------------------------------------------------------

def test_the_file_name_is_percent_decoded_then_sanitised(client, ed):
    pv = staged(client, ed, book([{"de_name": "X"}]), name="../../etc/Modèle données\n.xlsx")
    assert pv["batch"]["file_name"] == "Modèle données.xlsx"
    assert bulk.clean_file_name("C%3A%5Cusers%5Cben%5Cplan.xlsx") == "plan.xlsx"
    assert bulk.clean_file_name("%E2%80%AEevil%00") == "evil"
    assert bulk.clean_file_name(None) == bulk.clean_file_name("  ") == "upload.xlsx"
    assert len(bulk.clean_file_name("a" * 400)) == 255


# ---- the no-commit forms leave the existing routes as they were ------------------------------------

def test_the_existing_entity_routes_still_commit(client, ed, db):
    de = entity(client, ed, "Customer")
    r = client.patch(f"{API}/data-entities/{de['data_entity_id']}", headers=ed["h"],
                     json={"row_version": de["row_version"], "description": "x"})
    assert r.status_code == 200 and live(db, ed["p"])["DE-0001"].description == "x"
