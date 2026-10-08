"""The one spreadsheet reader and writer of the import pipeline (design §7.12, slice 4a spec §4, §7).

Reading an upload is three gates, in order, all before a business table is touched:
1. `zip_prepass` decompresses every member in chunks and counts real bytes (100 MB in all,
   20 MB for sharedStrings.xml), and refuses any XML part with a document type declaration or a
   non-UTF-8 encoding — so a DTD or an external entity is refused, never resolved.
2. openpyxl parses read-only, with defusedxml in place (checked at import below), so entity
   declarations are refused a second time by the parser itself.
3. `read_rows` streams the one sheet carrying the target's code (4a-R11) and stops at the
   5,000-row / 200-column caps while streaming, never after materialising the sheet.
Nothing is written to disk: the body is bytes in memory, and openpyxl reads it from BytesIO.

Writing goes through ONE `write_cell` for every cell, the hidden lookup sheet included (4a-R9):
text is forced to the string type, and a value starting with = + - @ tab or CR gets Excel's
quotePrefix style, so it is shown and re-read as text with no apostrophe added to its content.
Row 1 is hidden and locked: `sheet:<TARGET>` in A1, the column codes from B1 (4a-R5). Columns
are matched by those codes, never by position or by the translated header in row 2."""
import io
import re
import threading
import zipfile
from contextlib import contextmanager
from dataclasses import dataclass, field

import openpyxl
from defusedxml import DefusedXmlException
from fastapi import HTTPException
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet._reader import WorkSheetParser
from openpyxl.xml.constants import SHEET_MAIN_NS

# The parser's own hardening is the second gate; never run without it (criterion 75).
if not openpyxl.DEFUSEDXML or openpyxl.LXML:
    raise RuntimeError("openpyxl must parse through defusedxml (install defusedxml, not lxml)")

MB = 1024 * 1024
MAX_DECOMPRESSED_BYTES = 100 * MB           # the zip-bomb cap (§7.12)
MAX_SHARED_STRINGS_BYTES = 20 * MB
MAX_MEMBERS = 1_000                         # a real workbook has a few dozen parts
MAX_ROWS, MAX_COLS = 5_000, 200             # data rows; columns (§7.12)
EXCEL_MAX_ROW = 1_048_576
CHUNK = 64 * 1024
SHEET_TAG = "sheet:"                        # A1 of hidden row 1 (4a-R11)
FIRST_DATA_ROW = 3                          # row 1 codes (hidden), row 2 headers
SHEET_NAME_CAP = 31                         # Excel's own limit; untrusted text in every message
INLINE_LIST_CAP = 255                       # Excel's inline data-validation limit
LOOKUP_SHEET = "_lists"
ESCAPE_LEADS = ("=", "+", "-", "@", "\t", "\r")      # §7.12 escape_cell, the full set (4a-R9)
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_XML_PART = re.compile(r"\.(xml|rels|vml)$", re.IGNORECASE)
_DECLARED_ENCODING = re.compile(rb"""encoding\s*=\s*["']([A-Za-z0-9._-]+)["']""")
_FORMULA_TAG = f"{{{SHEET_MAIN_NS}}}f"
# clean_cell (4a-R15): str.strip() covers Unicode whitespace, NBSP included; these invisible
# format characters are not whitespace to Python but are to a consultant.
_INVISIBLE_EDGES = "​‌‍⁠﻿"


def refuse(message: str, status: int = 422) -> HTTPException:
    return HTTPException(status, message)


def sheet_label(name: str) -> str:
    """A sheet name as it may appear in a message: untrusted, cut to 31 characters."""
    return "'" + str(name)[:SHEET_NAME_CAP] + "'"


# ---- the parse semaphore -----------------------------------------------------------------

_PARSE_SLOTS = threading.BoundedSemaphore(2)


@contextmanager
def parse_slot():
    """At most two uploads parse at once (per process); a third is told to retry (429)."""
    if not _PARSE_SLOTS.acquire(blocking=False):
        raise refuse("Two uploads are already being checked. Try again in a moment.", 429)
    try:
        yield
    finally:
        _PARSE_SLOTS.release()


# ---- gate 1: the zip pre-pass ------------------------------------------------------------

def _check_xml_head(name: str, head: bytes) -> None:
    """An XML part must be UTF-8 (all Excel writes), so the DTD scan below sees plain bytes:
    a UTF-16 or EBCDIC part could otherwise spell <!DOCTYPE in bytes the scan cannot read."""
    if head.startswith((b"\xff\xfe", b"\xfe\xff")) or b"\x00" in head[:4]:
        raise refuse(f"The file is not a valid workbook: part {name[:100]} is not UTF-8.")
    m = _DECLARED_ENCODING.search(head[:200])
    if m and m.group(1).lower().replace(b"_", b"-") not in (b"utf-8", b"utf8"):
        raise refuse(f"The file is not a valid workbook: part {name[:100]} is not UTF-8.")


_OVERRIDE = re.compile(rb"<Override\b[^>]*>", re.IGNORECASE)
_PART_NAME = re.compile(rb"""PartName\s*=\s*["']/?([^"']+)["']""")


def _looks_xml(head: bytes) -> bool:
    """XML by content, not by name: a part can live at any path the package names. A UTF-16
    part counts too, so _check_xml_head refuses it."""
    h = head.removeprefix(b"\xef\xbb\xbf").removeprefix(b"\xff\xfe").removeprefix(b"\xfe\xff")
    h = h.lstrip(b" \t\r\n")
    return h[:1] == b"<" or h[:2] == b"\x00<"


def _shared_string_parts(archive: zipfile.ZipFile) -> set[str]:
    """The parts [Content_Types].xml declares as shared strings (openpyxl finds them that way)."""
    try:
        info = archive.getinfo("[Content_Types].xml")
    except KeyError:
        return set()
    if info.file_size > MB:
        raise refuse("The file is not a valid workbook: its content types are too large.")
    with archive.open(info) as f:
        types = f.read(MB + 1)
    if len(types) > MB:
        raise refuse("The file is not a valid workbook: its content types are too large.")
    out = set()
    for el in _OVERRIDE.findall(types):
        m = _PART_NAME.search(el)
        if m and b"sharedstrings+xml" in el.lower():
            out.add(m.group(1).decode("utf-8", "replace").lower())
    return out


def zip_prepass(body: bytes) -> None:
    """Decompress every member in chunks, counting real bytes (never the header's claim), and
    refuse a document type declaration in any XML part before any parser sees it."""
    try:
        archive = zipfile.ZipFile(io.BytesIO(body))
    except (zipfile.BadZipFile, ValueError, OSError):
        raise refuse("The file is not an .xlsx workbook.")
    with archive:
        members = archive.infolist()
        if len(members) > MAX_MEMBERS:
            raise refuse("The workbook has too many parts.")
        try:
            strings_parts = _shared_string_parts(archive)
        except HTTPException:
            raise
        except Exception:
            raise refuse("The file is not a readable .xlsx workbook.")
        total = 0
        for info in members:
            is_xml = bool(_XML_PART.search(info.filename))
            name = info.filename.lower()
            is_strings = name.endswith("sharedstrings.xml") or name in strings_parts
            seen, tail, first = 0, b"", True
            try:
                with archive.open(info) as member:
                    while chunk := member.read(CHUNK):
                        seen += len(chunk)
                        total += len(chunk)
                        if total > MAX_DECOMPRESSED_BYTES:
                            raise refuse("The workbook expands to more than 100 MB.")
                        if is_strings and seen > MAX_SHARED_STRINGS_BYTES:
                            raise refuse("The workbook's text table expands to more than 20 MB.")
                        if first:
                            is_xml = is_xml or _looks_xml(chunk)
                            if is_xml:
                                _check_xml_head(info.filename, chunk)
                            first = False
                        if is_xml:
                            window = (tail + chunk).lower()
                            if b"<!doctype" in window or b"<!entity" in window:
                                raise refuse("The workbook contains an XML document type "
                                             "declaration, which is not accepted.")
                            tail = window[-16:]
            except HTTPException:
                raise
            except Exception:            # bad CRC, encryption, unknown compression, truncation
                raise refuse("The file is not a readable .xlsx workbook.")


# ---- gate 2 + 3: streaming the target's sheet ------------------------------------------------

class _Marker:
    def __init__(self, name: str) -> None:
        self.name = name

    def __repr__(self) -> str:
        return self.name


NO_CACHED_VALUE = _Marker("NO_CACHED_VALUE")     # a formula Excel never calculated
CELL_ERROR = _Marker("CELL_ERROR")               # #N/A, #REF! …


class _Parser(WorkSheetParser):
    """openpyxl's streaming row parser (data_only: a formula reads as its cached value), told
    apart from a blank: a formula with no cached value, or an error value, is marked."""

    def parse_cell(self, element):
        cell = super().parse_cell(element)
        if cell["data_type"] == "e":
            cell["value"] = CELL_ERROR
        elif cell["value"] is None and element.find(_FORMULA_TAG) is not None:
            cell["value"] = NO_CACHED_VALUE
        return cell


def clean_cell(value):
    """Every text cell, before validation or lookup (4a-R15): Unicode whitespace (NBSP
    included) and invisible format characters stripped at both ends; empty → None."""
    if isinstance(value, str):
        value = value.strip().strip(_INVISIBLE_EDGES).strip()
        return value or None
    return value


@dataclass
class SheetRead:
    sheet_name: str
    rows: list[tuple[int, dict]]                 # (sheet row number, {column code: value})
    ignored_columns: list[str] = field(default_factory=list)


def _rows(wb, ws):
    src = ws._get_source()
    parser = _Parser(src, ws._shared_strings, data_only=True, epoch=wb.epoch,
                     date_formats=wb._date_formats, timedelta_formats=wb._timedelta_formats)
    try:
        for idx, cells in parser.parse():
            if idx > EXCEL_MAX_ROW:
                raise refuse("The sheet has a row beyond Excel's last row.")
            yield idx, cells
    finally:
        src.close()


def _first_cell(wb, ws):
    for idx, cells in _rows(wb, ws):
        if idx != 1:
            return None
        a1 = next((c["value"] for c in cells if c["column"] == 1), None)
        return clean_cell(a1)
    return None


def _unreadable(exc: BaseException) -> HTTPException:
    """openpyxl wraps a parser refusal in a ValueError; name the entity refusal if it is one."""
    seen = 0
    while exc is not None and seen < 10:
        if isinstance(exc, DefusedXmlException):
            return refuse("The workbook contains an XML entity declaration, which is not accepted.")
        exc, seen = exc.__cause__ or exc.__context__, seen + 1
    return refuse("The file is not a readable .xlsx workbook.")


def _open(body: bytes):
    try:
        return openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True,
                                      keep_links=False, keep_vba=False)
    except HTTPException:
        raise
    except Exception as e:
        raise _unreadable(e)


def read_rows(body: bytes, target_code: str, column_codes: tuple[str, ...]) -> SheetRead:
    """The data rows of the one sheet tagged `sheet:<target_code>`, keyed by column code.
    Blank rows are skipped; text is cleaned. File-level problems are a 422 with nothing read:
    no tagged sheet or several, a mapped column missing or repeated, the caps."""
    zip_prepass(body)
    wb = _open(body)
    try:
        try:
            tag = SHEET_TAG + target_code
            sheets = wb.worksheets
            tagged = [ws for ws in sheets if _first_cell(wb, ws) == tag]
            if len(tagged) != 1:
                names = ", ".join(sheet_label(ws.title) for ws in (tagged or sheets)) or "none"
                if not tagged:
                    raise refuse(f"No sheet is marked {tag} (row 1, cell A1). Sheets: {names}. "
                                 "Start from a VB template or export.")
                raise refuse(f"More than one sheet is marked {tag}: {names}. Keep one.")
            ws = tagged[0]
            return _read_sheet(wb, ws, column_codes)
        except HTTPException:
            raise
        except Exception as e:
            raise _unreadable(e)
    finally:
        wb.close()


def _read_sheet(wb, ws, column_codes: tuple[str, ...]) -> SheetRead:
    by_col: dict[int, str] = {}
    ignored: list[str] = []
    rows: list[tuple[int, dict]] = []
    last = 0
    for idx, cells in _rows(wb, ws):
        if idx <= last:                     # row numbers come from the file: never trust them
            raise refuse("The sheet's rows are out of order or repeated.")
        last = idx
        if any(c["column"] > MAX_COLS and c["value"] is not None for c in cells):
            raise refuse(f"The sheet has more than {MAX_COLS} columns.")
        if idx == 1:
            for c in cells:
                code = clean_cell(c["value"]) if c["column"] > 1 else None
                if code is None:
                    continue
                code = str(code)
                if code in by_col.values():
                    raise refuse(f"Column code {code[:40]} appears twice in row 1.")
                if code in column_codes:
                    by_col[c["column"]] = code
                else:
                    ignored.append(code[:40])
            missing = [c for c in column_codes if c not in by_col.values()]
            if missing:
                raise refuse("The sheet is missing these columns: " + ", ".join(missing) +
                             ". Start from a VB template or export.")
            continue
        if idx < FIRST_DATA_ROW:
            continue
        values = {code: None for code in column_codes}
        for c in cells:
            code = by_col.get(c["column"])
            if code is not None:
                values[code] = clean_cell(c["value"])
        if all(v is None for v in values.values()):
            continue
        rows.append((idx, values))
        if len(rows) > MAX_ROWS:
            raise refuse(f"The sheet has more than {MAX_ROWS:,} rows.")
    if not by_col:
        raise refuse("The sheet has no column codes in row 1. Start from a VB template or export.")
    return SheetRead(ws.title, rows, ignored)


# ---- writing ---------------------------------------------------------------------------------

@dataclass(frozen=True)
class OutColumn:
    code: str
    header: str
    width: int = 20
    locked: bool = False                         # a protected column (the row token)
    choices: tuple[str, ...] | None = None       # a data-validation list (resolved codes)


@dataclass(frozen=True)
class OutSheet:
    title: str
    target_code: str
    columns: tuple[OutColumn, ...]
    rows: list[list] = field(default_factory=list)   # values in column order


_UNLOCKED = Protection(locked=False)


def write_cell(ws, row: int, column: int, value, *, locked: bool = True, bold: bool = False):
    """THE one cell writer (4a-R9). Text is always the string type, never a formula; a value
    starting with = + - @ tab or CR is marked quotePrefix, which Excel shows as text without
    adding an apostrophe to the content, so the value reads back unchanged."""
    cell = ws.cell(row=row, column=column)
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)     # control characters XML cannot carry
        cell.value = value
        cell.data_type = "s"
        if value.startswith(ESCAPE_LEADS):
            cell.quotePrefix = True
    elif value is not None:
        cell.value = value
    if not locked:
        cell.protection = _UNLOCKED
    if bold:
        cell.font = Font(bold=True)
    return cell


def _inline_list(choices: tuple[str, ...]) -> str | None:
    """'"A,B,C"' when Excel can hold the list inline, else None (it goes on the lookup sheet)."""
    if any("," in c or '"' in c for c in choices):
        return None
    inline = '"' + ",".join(choices) + '"'
    return inline if len(inline) - 2 <= INLINE_LIST_CAP else None


def write_book(sheets: list[OutSheet]) -> bytes:
    """Template or export: one sheet per target, protected but editable where editing is meant,
    plus one hidden lookup sheet when a list is too long to sit inline."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    lookup, lookup_cols = None, 0
    for spec in sheets:
        ws = wb.create_sheet(spec.title[:SHEET_NAME_CAP])
        write_cell(ws, 1, 1, SHEET_TAG + spec.target_code)
        ws.column_dimensions["A"].hidden = True
        last_row = FIRST_DATA_ROW + MAX_ROWS - 1
        for i, col in enumerate(spec.columns, start=2):
            letter = get_column_letter(i)
            write_cell(ws, 1, i, col.code)
            write_cell(ws, 2, i, col.header, bold=True)
            dim = ws.column_dimensions[letter]
            dim.width = col.width
            if not col.locked:
                dim.protection = _UNLOCKED            # cells typed later stay editable
            if col.choices:
                formula = _inline_list(col.choices)
                if formula is None:
                    if lookup is None:
                        lookup = wb.create_sheet(LOOKUP_SHEET)
                        lookup.sheet_state = "hidden"
                    lookup_cols += 1
                    write_cell(lookup, 1, lookup_cols, col.code)
                    for r, choice in enumerate(col.choices, start=2):
                        write_cell(lookup, r, lookup_cols, choice)
                    lc = get_column_letter(lookup_cols)
                    formula = f"'{LOOKUP_SHEET}'!${lc}$2:${lc}${len(col.choices) + 1}"
                dv = DataValidation(type="list", formula1=formula, allow_blank=True)
                dv.add(f"{letter}{FIRST_DATA_ROW}:{letter}{last_row}")
                ws.add_data_validation(dv)
        for r, values in enumerate(spec.rows, start=FIRST_DATA_ROW):
            for i, (col, value) in enumerate(zip(spec.columns, values), start=2):
                write_cell(ws, r, i, value, locked=col.locked)
        ws.row_dimensions[1].hidden = True
        ws.freeze_panes = "B3"
        # Protection keeps hands off row 1 and the token column; it is a guard rail, not
        # security (the token's signature is). Sorting, filtering and adding rows stay allowed.
        p = ws.protection
        p.sheet = True
        p.sort = p.autoFilter = p.insertRows = p.deleteRows = False
        p.formatCells = p.formatColumns = p.formatRows = False
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
