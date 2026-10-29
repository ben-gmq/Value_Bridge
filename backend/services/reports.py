"""Slice 1f client reports (docs/slice1f_spec.md §3). Read-only and derived on every read from live rows. The caller
owns access: every route is project_ctx(VIEW).

A part that needs a component the project has switched off is None, never an empty list — an empty list would read
"no gaps" (spec A1). The one thing stored is the REPORT_PRINTED audit row a print writes (§3.6)."""
import io
import re
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo

import openpyxl
from fastapi import HTTPException
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from models import (AppUser, Application, ApplicationBfcNode, ApplicationSolution, BfcNode, BrDataEntity,
                    BrFunctionRequirement, BrSolution, BrSystemFlow, BusinessRequirement, DataEntity, FunctionRequirement,
                    Project, Solution)
from models.shared import FR_BEHAVIOURS
from services import (area_match, audit, br_flows, components, function_requirement as fr_service,
                      interface as iface_service, xlsx)
from services.code_master import resolve

BEHAVIOUR_ORDER = ("FORM", "BATCH", "INTERFACE", "REPORT", "OTHER")
assert set(BEHAVIOUR_ORDER) == set(FR_BEHAVIOURS)          # a new behaviour gets its place here (sara 1f L1)


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


# ---- the stamp every printed or exported report carries (D-42) ----------------------------------------------------
ICT = ZoneInfo("Asia/Bangkok")


def stamp_time(at: datetime) -> str:
    """The one way a report states when it was worked out: ICT, as text with its zone (D8). The Excel stamp and the
    print stamp both use it, so the two never drift (VB-005 piece-4 L2)."""
    return at.astimezone(ICT).strftime("%Y-%m-%d %H:%M ICT")


def _hier_key(code: str) -> tuple:
    return tuple(int(p) if p.isdigit() else p for p in code.split("."))


class _Scope:
    """The live rows every report reads, loaded once per request. A BR counts only on a live to-be step; a link
    counts only when both of its ends are live."""

    def __init__(self, db: Session, project_id: int):
        self.on = components.enabled(db, project_id)
        self.steps = {n.bfc_node_id: n for n in db.scalars(select(BfcNode).where(
            BfcNode.project_id == project_id, BfcNode.is_active, BfcNode.bfc_set_code == "TO_BE"))}
        self.brs = {b.br_id: b for b in db.scalars(select(BusinessRequirement).where(
            BusinessRequirement.project_id == project_id, BusinessRequirement.is_active))
            if b.bfc_node_id in self.steps}
        self.frs = {f.fr_id: f for f in db.scalars(select(FunctionRequirement).where(
            FunctionRequirement.project_id == project_id, FunctionRequirement.is_active))}
        self.sols = {s.solution_id: s for s in db.scalars(select(Solution).where(
            Solution.project_id == project_id, Solution.is_active))}
        self.br_fr = [(br, fr) for br, fr in db.execute(select(BrFunctionRequirement.br_id, BrFunctionRequirement.fr_id)
                      .where(BrFunctionRequirement.project_id == project_id, BrFunctionRequirement.is_active))
                      if br in self.brs and fr in self.frs]
        self.br_sol = [(br, s) for br, s in db.execute(select(BrSolution.br_id, BrSolution.solution_id).where(
                       BrSolution.project_id == project_id, BrSolution.is_active))
                       if br in self.brs and s in self.sols]
        self.apps = apps = {a.application_id: a for a in db.scalars(select(Application).where(
            Application.project_id == project_id, Application.is_active))}
        self.apps_of_sol: dict[int, list[dict]] = {}
        self.app_ids_of_sol: dict[int, set[int]] = {}
        for app_id, sol_id in db.execute(select(ApplicationSolution.application_id, ApplicationSolution.solution_id)
                                         .where(ApplicationSolution.project_id == project_id,
                                                ApplicationSolution.is_active)):
            if app_id in apps:
                a = apps[app_id]
                self.app_ids_of_sol.setdefault(sol_id, set()).add(app_id)
                self.apps_of_sol.setdefault(sol_id, []).append(
                    {"application_code": a.application_code, "application_name": a.application_name})
        for v in self.apps_of_sol.values():
            v.sort(key=lambda a: a["application_code"])

    def live_sol(self, sol_id: int | None) -> Solution | None:
        return self.sols.get(sol_id) if sol_id is not None else None

    def answers(self) -> set[tuple[int, int]]:
        """(BR, solution) pairs where the solution is a live, non-REJECTED answer of the BR (S5-2, D4)."""
        return {(b, s) for b, s in self.br_sol if self.sols[s].status_code != "REJECTED"}


def _sol_out(s: Solution | None) -> dict | None:
    if s is None:
        return None
    return {"solution_id": s.solution_id, "solution_number": s.solution_number, "solution_name": s.solution_name,
            "category_code": s.category_code, "status_code": s.status_code}


def _br_out(b: BusinessRequirement) -> dict:
    return {"br_id": b.br_id, "br_number": b.br_number, "br_statement": b.br_statement}


def _fr_out(f: FunctionRequirement) -> dict:
    return {"fr_id": f.fr_id, "fr_number": f.fr_number, "fr_name": f.fr_name, "fulfilment_status_code": f.fulfilment_status_code}


# ---- 3.1 traceability -------------------------------------------------------------------------------------------
def traceability(db: Session, project_id: int) -> dict:
    """One row per live (to-be step × BR on it × BR↔FR link × FR), grouped by FR type in behaviour order (criterion
    10a: an FR on three BRs is three rows; D-52: an FR reached only through a solution is none). Then the untraced
    FRs, and the non-technology solutions that answer a BR with no FR (criterion 10, D5)."""
    sc = _Scope(db, project_id)
    if "FUNCTION_REQUIREMENT" not in sc.on:
        return {"generated_at": _now(), "groups": None, "untraced": None, "answered_without_software": None}
    sol_on, app_on = "SOLUTION" in sc.on, "APPLICATION" in sc.on
    answers = sc.answers()

    def tag_part(f: FunctionRequirement, br_id: int | None) -> dict:
        if not sol_on:
            return {"solution": None, "applications": None, "tag_answers_br": None, "rejected_solution": None}
        s = sc.live_sol(f.solution_id)
        return {"solution": _sol_out(s),
                "applications": (sc.apps_of_sol.get(s.solution_id, []) if s else []) if app_on else None,
                "tag_answers_br": None if s is None or br_id is None else (br_id, s.solution_id) in answers,
                "rejected_solution": s is not None and s.status_code == "REJECTED"}

    groups: dict[int, dict] = {}
    for br_id, fr_id in sc.br_fr:
        f, b = sc.frs[fr_id], sc.brs[br_id]
        step = sc.steps[b.bfc_node_id]
        t = f.fr_type_code_row
        g = groups.setdefault(f.fr_type_code_id, {"fr_type_code": t.code, "fr_type_label": t.label,
                                                  "behaviour": f.fr_type_behaviour, "rows": []})
        g["rows"].append({"hier_code": step.hier_code, "step_name": step.node_name, **_br_out(b), **_fr_out(f),
                          **tag_part(f, br_id)})
    for g in groups.values():
        g["rows"].sort(key=lambda r: (_hier_key(r["hier_code"]), r["br_number"], r["fr_number"]))

    def order(g):
        b = g["behaviour"]
        return (BEHAVIOUR_ORDER.index(b) if b in BEHAVIOUR_ORDER else len(BEHAVIOUR_ORDER), g["fr_type_label"])
    linked = {fr for _, fr in sc.br_fr}
    untraced = sorted(({**_fr_out(f), **tag_part(f, None)} for f in sc.frs.values() if f.fr_id not in linked),
                      key=lambda r: r["fr_number"])
    without = None
    if sol_on:
        with_fr = {f.solution_id for f in sc.frs.values() if f.solution_id is not None}
        brs_of: dict[int, list[str]] = {}
        for b, s in answers:
            brs_of.setdefault(s, []).append(sc.brs[b].br_number)
        without = sorted(({**_sol_out(sc.sols[s]), "br_numbers": sorted(nums)} for s, nums in brs_of.items()
                          if s not in with_fr and sc.sols[s].category_code not in fr_service.TECH_CATEGORIES),
                         key=lambda r: r["solution_number"])
    return {"generated_at": _now(), "groups": sorted(groups.values(), key=order), "untraced": untraced,
            "answered_without_software": without}


# ---- 3.2 coverage -------------------------------------------------------------------------------------------------
def _line(code: str, severity: str, rows: list | None) -> dict:
    return {"code": code, "severity": severity, "rows": rows}


def coverage(db: Session, project_id: int) -> dict:
    """§4.3 br_coverage_report: one line per gap kind, each {code, severity, rows}, in the spec's §3.2 order. A line
    whose component is off has rows None (A1). Rows carry ids so the screen links each one to its item."""
    sc = _Scope(db, project_id)
    fr_on, sol_on, app_on = ("FUNCTION_REQUIREMENT" in sc.on, "SOLUTION" in sc.on, "APPLICATION" in sc.on)
    both = fr_on and sol_on
    brs = sorted(sc.brs.values(), key=lambda b: b.br_number)
    frs = sorted(sc.frs.values(), key=lambda f: f.fr_number)
    answers = sc.answers()
    des = {d.data_entity_id: d for d in db.scalars(select(DataEntity).where(
        DataEntity.project_id == project_id, DataEntity.is_active))}
    de_rows = [(br, de, crud) for br, de, crud in db.execute(
        select(BrDataEntity.br_id, BrDataEntity.data_entity_id, BrDataEntity.direction).where(
            BrDataEntity.project_id == project_id, BrDataEntity.is_active))
        if br in sc.brs and de in des]
    with_output = {br for br, _, direction in de_rows if direction == "O"}
    referenced = {de for _, de, _ in de_rows}
    linked_brs = {b for b, _ in sc.br_fr}
    linked_frs = {f for _, f in sc.br_fr}
    answered = {b for b, _ in answers}

    def br_rows(keep) -> list[dict]:
        return [_br_out(b) for b in brs if keep(b)]

    lines = [_line("br_without_output_de", "information", br_rows(lambda b: b.br_id not in with_output)),
             _line("br_without_solution", "information", br_rows(lambda b: b.br_id not in answered) if sol_on else None),
             _line("br_without_fr", "information", br_rows(lambda b: b.br_id not in linked_brs) if fr_on else None),
             _line("fr_without_br", "warning", [_fr_out(f) for f in frs if f.fr_id not in linked_frs] if fr_on else None),
             _line("fr_without_solution", "information",
                   [_fr_out(f) for f in frs if sc.live_sol(f.solution_id) is None] if both else None)]
    outside = None
    if both:     # a retired tag is fr_without_solution's row, not this one too (sara 1f-2 L3; the freeze line stays, N3)
        outside = [{**_fr_out(sc.frs[r["fr_id"]]), "solution_number": r["solution_number"]}
                   for r in fr_service.fr_solution_outside_its_brs(db, project_id)
                   if r["fr_id"] in sc.frs and sc.live_sol(sc.frs[r["fr_id"]].solution_id)]
    lines.append(_line("fr_solution_outside_its_brs", "warning", outside))
    reach = None
    if both:
        pairs = set(sc.br_fr)
        reach = sorted(({**_fr_out(f), **_br_out(sc.brs[b]), "solution_number": sc.sols[s].solution_number}
                        for f in frs if f.fr_id in linked_frs for b, s in answers
                        if s == f.solution_id and (b, f.fr_id) not in pairs),
                       key=lambda r: (r["fr_number"], r["br_number"]))
    lines.append(_line("solution_reach_without_link", "information", reach))
    flows = None
    if app_on:
        flows = []
        live = {fl.br_system_flow_id: fl for fl in db.scalars(select(BrSystemFlow).where(
            BrSystemFlow.project_id == project_id, BrSystemFlow.is_active))}
        with_flows = {fl.br_id for fl in live.values()}
        for b in brs:
            if b.br_id not in with_flows:
                continue
            for fid in sorted(fid for fid, covered_by in br_flows.coverage(db, b).items() if not covered_by):
                fl = live[fid]
                src, tgt = sc.apps.get(fl.source_application_id), sc.apps.get(fl.target_application_id)
                de = des.get(fl.data_entity_id)
                flows.append({**_br_out(b), "br_system_flow_id": fl.br_system_flow_id,
                              "de_number": de.de_number if de else None, "de_name": de.de_name if de else None,
                              "source_code": src.application_code if src else None,
                              "target_code": tgt.application_code if tgt else None})
    lines.append(_line("flow_uncovered", "warning" if fr_on else "information", flows))
    lines.append(_line("fr_rejected_tag", "warning",
                       [{**_fr_out(f), "solution_number": s.solution_number} for f in frs
                        if (s := sc.live_sol(f.solution_id)) is not None and s.status_code == "REJECTED"]
                       if both else None))
    tech = None
    if both:
        with_fr = {f.solution_id for f in frs if f.solution_id is not None}
        tech = sorted((_sol_out(sc.sols[s]) for s in {s for _, s in answers}
                       if sc.sols[s].category_code in fr_service.TECH_CATEGORIES and s not in with_fr),
                      key=lambda r: r["solution_number"])
    lines.append(_line("tech_solution_without_fr", "warning", tech))
    lines.append(_line("de_unreferenced", "information",
                       [{"data_entity_id": d.data_entity_id, "de_number": d.de_number, "de_name": d.de_name}
                        for d in sorted(des.values(), key=lambda d: d.de_number) if d.data_entity_id not in referenced]))
    return {"generated_at": _now(), "lines": lines}


# ---- 3.3 solution matrix ---------------------------------------------------------------------------------------
def solution_matrix(db: Session, project_id: int) -> dict:
    """§7.9 solution_matrix, criterion 10b: one cell per (to-be L1 area × solution category) holding the distinct live
    non-REJECTED solutions that answer a live BR under that area and the count of those BRs. Every category is a
    column, in SOLUTION_CATEGORY order, even when empty. Solutions on no live BR get their own row; REJECTED ones are
    only counted (A6). Needs SOLUTION."""
    sc = _Scope(db, project_id)
    if "SOLUTION" not in sc.on:
        return {"generated_at": _now(), "categories": None, "rows": None, "unlinked": None, "totals": None,
                "live_count": None, "rejected_count": None, "technology_share": None}
    cats = [{"code": c.code, "label": c.label} for c in resolve(db, project_id, "SOLUTION_CATEGORY")]
    areas = {n.hier_code: n for n in sc.steps.values() if n.level_no == 1 and not n.is_process}
    live = {sid: s for sid, s in sc.sols.items() if s.status_code != "REJECTED"}
    cells: dict[str, dict[str, tuple[set, set]]] = {}
    linked: set[int] = set()
    for br_id, sid in sc.br_sol:
        if sid not in live:
            continue
        linked.add(sid)
        area = sc.steps[sc.brs[br_id].bfc_node_id].hier_code.split(".", 1)[0]
        sols, brs = cells.setdefault(area, {}).setdefault(live[sid].category_code, (set(), set()))
        sols.add(sid)
        brs.add(br_id)

    def row(by_cat: dict) -> dict:
        empty = (set(), set())
        return {"cells": {c["code"]: {"solutions": len(by_cat.get(c["code"], empty)[0]),
                                      "brs": len(by_cat.get(c["code"], empty)[1])} for c in cats},
                "total": {"solutions": len(set().union(*(v[0] for v in by_cat.values()))),
                          "brs": len(set().union(*(v[1] for v in by_cat.values())))}}

    rows = [{"bfc_node_id": a.bfc_node_id, "hier_code": code, "area_name": a.node_name, **row(cells.get(code, {}))}
            for code, a in sorted(areas.items(), key=lambda kv: _hier_key(kv[0]))]
    unlinked: dict[str, tuple[set, set]] = {}
    for sid, s in live.items():
        if sid not in linked:
            unlinked.setdefault(s.category_code, (set(), set()))[0].add(sid)
    totals = {c["code"]: sum(1 for s in live.values() if s.category_code == c["code"]) for c in cats}
    share = round(totals.get(fr_service.TECHNOLOGY, 0) * 100 / len(live), 1) if live else None
    return {"generated_at": _now(), "categories": cats, "rows": rows, "unlinked": row(unlinked)["cells"],
            "totals": totals, "live_count": len(live), "rejected_count": len(sc.sols) - len(live),
            "technology_share": share}


# ---- §3.4 the system gap (function 67; D-8, RP-1) -----------------------------------------------------------------
GAP_NEEDS = ("APPLICATION", "AS_IS", "SOLUTION")         # to-be allocation joins BR_SOLUTION and APPLICATION_SOLUTION (D2)
SUPPORTING = ("AS_IS", "BOTH")                            # lifecycles that count as today's support (S6b-5)


def _step_out(n: BfcNode) -> dict:
    return {"bfc_node_id": n.bfc_node_id, "hier_code": n.hier_code, "node_name": n.node_name}


def _app_out(a: Application) -> dict:
    return {"application_id": a.application_id, "application_code": a.application_code,
            "application_name": a.application_name}


def application_gap(db: Session, project_id: int, level: int = 1) -> dict:
    """The migration backlog, per system inside each matched area (RP-1). As-is support = a live support link on a live
    as-is process to a live AS_IS / BOTH application; to-be allocation = a live to-be process ⋈ live BR ⋈ a live,
    non-REJECTED answer (as coverage counts answers, S5-2) ⋈ a live allocation of a live application. Areas are paired
    by area_match, the volume summary's rule. The input counts come back so an empty project never reads "no gap"
    (D3); with APPLICATION, AS_IS or SOLUTION off every part is None (A1)."""
    if level not in (1, 2):
        raise HTTPException(422, "Level must be 1 or 2")
    sc = _Scope(db, project_id)
    if any(c not in sc.on for c in GAP_NEEDS):
        return {"generated_at": _now(), "level": level, "counts": None, "unallocated": None, "unsupported": None,
                "retiring": None, "as_is_support_by_to_be": None, "unmatched": None}
    as_is = {n.bfc_node_id: n for n in db.scalars(select(BfcNode).where(
        BfcNode.project_id == project_id, BfcNode.is_active, BfcNode.bfc_set_code == "AS_IS"))}
    support: dict[int, set[int]] = {}                    # as-is step → applications supporting it today
    by_to_be: list[tuple[BfcNode, Application]] = []
    links = 0
    for step_id, app_id in db.execute(select(ApplicationBfcNode.bfc_node_id, ApplicationBfcNode.application_id).where(
            ApplicationBfcNode.project_id == project_id, ApplicationBfcNode.is_active)):
        step, app = as_is.get(step_id), sc.apps.get(app_id)
        if step is None or not step.is_process or app is None:
            continue
        if app.lifecycle_code in SUPPORTING:
            support.setdefault(step_id, set()).add(app_id)
            links += 1
        else:
            by_to_be.append((step, app))                 # information, never today's support (S6b-5)
    alloc: dict[int, set[int]] = {}                      # to-be step → applications allocated to it
    for br_id, sol_id in sc.answers():
        apps = sc.app_ids_of_sol.get(sol_id)
        if apps:
            alloc.setdefault(sc.brs[br_id].bfc_node_id, set()).update(apps)
    areas = {"AS_IS": {c: n for c, n in ((n.hier_code, n) for n in as_is.values())
                       if n.level_no == level and not n.is_process},
             "TO_BE": {n.hier_code: n for n in sc.steps.values() if n.level_no == level and not n.is_process}}
    pairs, unmatched = area_match.pair_areas(areas)
    under = lambda nodes, code: [n for n in nodes if n.is_process and n.hier_code.startswith(code + ".")]  # noqa: E731
    unallocated, unsupported = [], []
    for code, a_area, b_area in pairs:
        old_steps, new_steps = under(as_is.values(), code), under(sc.steps.values(), code)
        old_apps = set().union(*(support.get(s.bfc_node_id, set()) for s in old_steps))
        new_apps = set().union(*(alloc.get(s.bfc_node_id, set()) for s in new_steps))
        area = {"area_code": code, "area_name": b_area.node_name}
        for s in sorted(old_steps, key=lambda n: _hier_key(n.hier_code)):
            for app_id in sorted(support.get(s.bfc_node_id, set()) - new_apps, key=lambda i: sc.apps[i].application_code):
                unallocated.append({**area, "step": _step_out(s), "application": _app_out(sc.apps[app_id])})
        for s in sorted(new_steps, key=lambda n: _hier_key(n.hier_code)):
            for app_id in sorted(alloc.get(s.bfc_node_id, set()) - old_apps, key=lambda i: sc.apps[i].application_code):
                unsupported.append({**area, "step": _step_out(s), "application": _app_out(sc.apps[app_id])})
    allocated = set().union(*(sc.app_ids_of_sol.get(s, set()) for _, s in sc.answers()))
    retiring = sorted((sc.apps[i] for i in set().union(*support.values()) - allocated),
                      key=lambda a: a.application_code)                                         # A7
    return {"generated_at": _now(), "level": level,
            "counts": {"as_is_processes": sum(1 for n in as_is.values() if n.is_process), "support_links": links,
                       "to_be_allocations": sum(len(v) for v in alloc.values())},
            "unallocated": unallocated, "unsupported": unsupported, "retiring": [_app_out(a) for a in retiring],
            "as_is_support_by_to_be": [{"step": _step_out(s), "application": _app_out(a)}
                                       for s, a in sorted(by_to_be, key=lambda p: (_hier_key(p[0].hier_code),
                                                                                   p[1].application_code))],
            "unmatched": unmatched}


# ---- §3.5 the interface list as an Excel file (function 61) -------------------------------------------------------
IFACE_HEADERS = ("Interface", "Name", "Type", "Source → target", "Pattern", "Method", "Frequency", "Data entities carried",
                 "Solution", "BRs", "Status", "Effort (days)", "Complete")
IFACE_WIDTHS = (12, 40, 16, 22, 14, 12, 14, 44, 12, 22, 16, 12, 34)
MISSING_TEXT = {"endpoints": "endpoints", "method": "method", "data_entities": "data entities"}


def _app_label(db: Session, project_id: int, app_id: int | None) -> str:
    a = db.scalar(select(Application).where(Application.application_id == app_id,
                                            Application.project_id == project_id)) if app_id is not None else None
    return "any" if app_id is None else (a.application_code if a else f"#{app_id}")


def interface_workbook(db: Session, project_id: int, by: str, source: int | None = None, target: int | None = None,
                       status: str | None = None) -> bytes:
    """The rows of `interface_list` under the screen's filters, one sheet: the print stamp, the filters (and how many
    undescribed interfaces a pair filter left out, D6), a header row, one row per interface. Every value goes through
    xlsx.write_cell. A report, not an import template: no code row, no protection."""
    data = iface_service.interface_list(db, project_id, source, target, status)
    effort_on = "EFFORT" in components.enabled(db, project_id)
    project = db.get(Project, project_id)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Interfaces"
    w = xlsx.write_cell
    w(ws, 1, 1, f"{project.project_code} {project.project_name}", bold=True)
    w(ws, 2, 1, "LIVE — not a baseline")
    w(ws, 3, 1, f"Generated {stamp_time(datetime.now(timezone.utc))} by {by}")
    filtered = source is not None or target is not None or status is not None
    w(ws, 4, 1, "Filters: " + (f"source {_app_label(db, project_id, source)}, target {_app_label(db, project_id, target)}"
                               f", status {status or 'any'}" if filtered else "none"))
    row = 5
    if filtered and data["unfilterable_incomplete"]:
        w(ws, row, 1, f"{data['unfilterable_incomplete']} interfaces not yet described are not in this filtered list")
        row += 1
    row += 1
    for i, (h, width) in enumerate(zip(IFACE_HEADERS, IFACE_WIDTHS), start=1):
        w(ws, row, i, h, bold=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    for r in data["rows"]:
        row += 1
        route = f"{r['source_application_code']} → {r['target_application_code']}" if r["is_described"] else None
        carried = ", ".join(f"{c['de_number']} {c['de_name']}" + (f" ({c['note']})" if c["note"] else "")
                            for c in r["carried"])
        complete = "Complete" if r["is_complete"] else "Missing: " + ", ".join(MISSING_TEXT[m] for m in r["missing"])
        effort = float(r["effort_days"]) if effort_on and r["effort_days"] is not None else None
        for i, v in enumerate((r["interface_number"], r["interface_name"], r["fr_type_label"], route, r["pattern_label"],
                               r["method_label"], r["frequency_label"], carried or None, r["solution_number"],
                               ", ".join(r["br_numbers"]) or None, r["fulfilment_status_label"], effort, complete),
                              start=1):
            w(ws, row, i, v)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ---- §3.6 print to PDF (function 143): the stamp, and the one audit row a print writes -------------------------------
PRINTABLE = ("traceability", "coverage", "solution_matrix", "interfaces", "application_gap")
STAMP_WINDOW_MINUTES = 30
FILTERS_CAP = 200
_CONTROL = re.compile(r"[\x00-\x1f\x7f]")


def _checked_time(value: str) -> datetime:
    """Round 2 N2: the report's own generated_at, sent back by the browser — an ISO time with its zone, within the
    last 30 minutes, so a malformed or backdated value never reaches the append-only audit row."""
    try:
        at = datetime.fromisoformat(value)
    except (ValueError, TypeError):
        raise HTTPException(422, "generated_at is not a time") from None
    if at.tzinfo is None:
        raise HTTPException(422, "generated_at must carry its time zone")
    now = datetime.now(timezone.utc)
    if not now - timedelta(minutes=STAMP_WINDOW_MINUTES) <= at <= now + timedelta(seconds=5):
        raise HTTPException(422, "The report is out of date — refresh it, then print")
    return at.astimezone(timezone.utc)


def print_stamp(db: Session, actor: AppUser, project: Project, report: str, filters: str | None, generated_at: str,
                source_ip: str | None = None) -> dict:
    """D-42: what every printed page carries, and one REPORT_PRINTED row. Called once per press of Print / PDF, after
    the screen refetched the report; generated_at is that response's own time (P3)."""
    if report not in PRINTABLE:
        raise HTTPException(422, "That report cannot be printed")
    if filters is not None and (len(filters) > FILTERS_CAP or _CONTROL.search(filters)):
        raise HTTPException(422, f"Filters must be plain text of at most {FILTERS_CAP} characters")
    at = _checked_time(generated_at)
    audit.record(db, "REPORT_PRINTED", actor_id=actor.app_user_id, project_id=project.project_id,
                 target_table="project", target_id=project.project_id, source_ip=source_ip,
                 detail={"report": report, "filters": filters, "generated_at": at.isoformat()})
    db.commit()
    return {"report": report, "project_code": project.project_code, "project_name": project.project_name,
            "is_live": True, "generated_at": at.isoformat(), "generated_at_text": stamp_time(at),
            "generated_by": actor.display_name, "filters": filters}
